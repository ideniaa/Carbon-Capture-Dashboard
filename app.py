from __future__ import annotations

import ast
import csv
import io
import json
import math
import os
import random
import subprocess
import sys
import threading
import time
from dataclasses import dataclass
from datetime import datetime, timedelta, timezone
from typing import Any

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
except ImportError as exc:
    raise ImportError(
        "openpyxl is required for XLSX exports. Install it with `pip install openpyxl`."
    ) from exc

from flask import Flask, flash, jsonify, redirect, render_template, request, send_file, url_for

try:
    import pyodbc
except ImportError:
    pyodbc = None

MODBUS_HOST = os.getenv("MODBUS_HOST", "").strip()
MODBUS_PORT = int(os.getenv("MODBUS_PORT", "502"))
MODBUS_UNIT_ID = int(os.getenv("MODBUS_UNIT_ID", "2"))

try:
    from pymodbus.client import ModbusTcpClient as _ModbusTcpClient
    _pymodbus_available = True
except ImportError:
    try:
        from pymodbus.client.sync import ModbusTcpClient as _ModbusTcpClient  # type: ignore[no-redef]
        _pymodbus_available = True
    except ImportError:
        _pymodbus_available = False

SEQUENCE_STATE_MAP: dict[int, str] = {
    0: "Ready", 1: "Preparing", 2: "Waiting", 3: "Running",
    4: "Idle", 5: "Stopping", 6: "Stopped", 7: "Halted",
    8: "Error", 9: "Blocked",
}

COMMAND_COIL_MANUAL_ADDRESS_MAP: dict[str, int] = {
    "Start Sequence": 1,
    "Stop Sequence": 2,
    "Break Sequence": 3,
    "Trigger Program A": 33,
    "Trigger Program B": 34,
    "Trigger Program X": 35,
}

COMMAND_COIL_MAP: dict[str, int] = {
    command: manual_address - 1
    for command, manual_address in COMMAND_COIL_MANUAL_ADDRESS_MAP.items()
}

app = Flask(__name__)
app.config["SECRET_KEY"] = os.getenv("SECRET_KEY", "change-me-in-production")

TITROLYZER_APP_PATH = os.getenv("TITROLYZER_APP_PATH", r"C:\Titrolyzer\Titrolyzer.py").strip()
TITROLYZER_PYTHON_EXE = os.getenv("TITROLYZER_PYTHON_EXE", sys.executable).strip()
TITROLYZER_PID_FILE = os.getenv("TITROLYZER_PID_FILE", os.path.join(app.root_path, "titrolyzer_app.pid")).strip()
TITROLYZER_STOPPED_FLAG = os.path.join(app.root_path, "titrolyzer_stopped.flag")


def load_titrolyzer_script_config(path: str) -> dict[str, str]:
    if not path or not os.path.exists(path):
        return {}
    wanted = {"SQL_SERVER", "SQL_DATABASE", "SQL_TABLE", "SQL_USERNAME", "SQL_PASSWORD"}
    config: dict[str, str] = {}
    try:
        with open(path, "r", encoding="utf-8", errors="replace") as handle:
            tree = ast.parse(handle.read(), filename=path)
        for node in tree.body:
            if not isinstance(node, ast.Assign):
                continue
            try:
                value = ast.literal_eval(node.value)
            except (ValueError, TypeError):
                continue
            for target in node.targets:
                if isinstance(target, ast.Name) and target.id in wanted:
                    config[target.id] = str(value).strip()
    except Exception as exc:
        app.logger.warning("Could not read Titrolyzer script config: %s", exc)
    return config


TITROLYZER_SCRIPT_CONFIG = load_titrolyzer_script_config(TITROLYZER_APP_PATH)


@dataclass
class DashboardState:
    sequence: str = "Idle"
    program_a: str = "Standby"
    program_b: str = "Standby"
    analyzer_on: bool | None = None
    analyzer_error: bool | None = None
    analyzer_failure: bool | None = None
    analyzer_timestamp: datetime | None = None
    connection_ok: bool = True


state = DashboardState()
history: list[dict[str, Any]] = []
chlorine_measurements: list[dict[str, Any]] = []
command_log: list[dict[str, str]] = []
results_entries: list[dict[str, Any]] = []
voltage_entries: list[dict[str, Any]] = []
announcements: list[dict[str, Any]] = []
_announcement_id_counter: int = 0
voltage_command_log: list[dict[str, str]] = []
_bootstrapped: bool = False
titrator_notes: dict[str, str] = {}
TITRATOR_NOTES_FILE = os.path.join(app.root_path, "titrator_notes.json")
voltage_notes: dict[str, str] = {}
VOLTAGE_NOTES_FILE = os.path.join(app.root_path, "voltage_notes.json")
COMMAND_LOG_FILE = os.path.join(app.root_path, "command_log.json")
VOLTAGE_COMMAND_LOG_FILE = os.path.join(app.root_path, "voltage_command_log.json")
SQL_SERVER_CONNECTION_STRING = os.getenv("SQL_SERVER_CONNECTION_STRING", "").strip()
SQL_SERVER = os.getenv("SQL_SERVER", TITROLYZER_SCRIPT_CONFIG.get("SQL_SERVER", "")).strip()
SQL_DATABASE = os.getenv("SQL_DATABASE", TITROLYZER_SCRIPT_CONFIG.get("SQL_DATABASE", "")).strip()
SQL_USERNAME = os.getenv("SQL_USERNAME", TITROLYZER_SCRIPT_CONFIG.get("SQL_USERNAME", "")).strip()
SQL_PASSWORD = os.getenv("SQL_PASSWORD", TITROLYZER_SCRIPT_CONFIG.get("SQL_PASSWORD", "")).strip()
SQL_DRIVER = os.getenv("SQL_DRIVER", "").strip()
SQL_TRUSTED_CONNECTION = os.getenv("SQL_TRUSTED_CONNECTION", "").strip().lower() in {"1", "true", "yes"}
SQL_SCHEMA = os.getenv("SQL_SCHEMA", "dbo").strip() or "dbo"

SQL_TITRATOR_TABLE = os.getenv("SQL_TITRATOR_TABLE", TITROLYZER_SCRIPT_CONFIG.get("SQL_TABLE", "Titrolyzer")).strip() or "Titrolyzer"
SQL_TITRATOR_RECORD_ID_COLUMN = os.getenv("SQL_TITRATOR_RECORD_ID_COLUMN", "RecordID").strip() or "RecordID"
SQL_TITRATOR_READING_TIME_COLUMN = os.getenv("SQL_TITRATOR_READING_TIME_COLUMN", "ReadingTime").strip() or "ReadingTime"
SQL_TITRATOR_TIMESTAMP_COLUMN = os.getenv("SQL_TITRATOR_TIMESTAMP_COLUMN", "Timestamp").strip() or "Timestamp"
SQL_TITRATOR_RESULT_ID_COLUMN = os.getenv("SQL_TITRATOR_RESULT_ID_COLUMN", "Result_ID").strip() or "Result_ID"
SQL_TITRATOR_VALUE_COLUMN = os.getenv("SQL_TITRATOR_VALUE_COLUMN", "Value").strip() or "Value"

SQL_CHLORINE_TABLE = os.getenv("SQL_CHLORINE_TABLE", "Chlorine").strip() or "Chlorine"
SQL_RESULTS_TABLE = os.getenv("SQL_RESULTS_TABLE", "Results").strip() or "Results"
SQL_VOLTAGE_TABLE = os.getenv("SQL_VOLTAGE_TABLE", "Voltage").strip() or "Voltage"
SQL_VOLTAGE_TABLE_PATTERN = os.getenv("SQL_VOLTAGE_TABLE_PATTERN", "V%").strip() or "V%"
SQL_VOLTAGE_TIMESTAMP_COLUMN = os.getenv("SQL_VOLTAGE_TIMESTAMP_COLUMN", "date_time").strip() or "date_time"
SQL_VOLTAGE_CH1_COLUMN = os.getenv("SQL_VOLTAGE_CH1_COLUMN", "channel_1").strip() or "channel_1"
SQL_VOLTAGE_CH2_COLUMN = os.getenv("SQL_VOLTAGE_CH2_COLUMN", "channel_2").strip() or "channel_2"
SQL_VOLTAGE_CH3_COLUMN = os.getenv("SQL_VOLTAGE_CH3_COLUMN", "channel_3").strip() or "channel_3"
SQL_VOLTAGE_CH4_COLUMN = os.getenv("SQL_VOLTAGE_CH4_COLUMN", "channel_4").strip() or "channel_4"
SQL_VOLTAGE_NOTE_COLUMN = os.getenv("SQL_VOLTAGE_NOTE_COLUMN", "").strip()
SQL_VOLTAGE_ROW_ID_COLUMN = os.getenv("SQL_VOLTAGE_ROW_ID_COLUMN", "row_num").strip()
SQL_WINCC_TABLES = os.getenv("SQL_WINCC_TABLES", "dbo.EQU_DB,dbo.EQU_DB2").strip()
SQL_WINCC_TABLE_NAMES = [name.strip() for name in SQL_WINCC_TABLES.split(",") if name.strip()] or ["dbo.EQU_DB", "dbo.EQU_DB2"]
SQL_WINCC_TIMESTAMP_COLUMN = os.getenv("SQL_WINCC_TIMESTAMP_COLUMN", "DateTime").strip() or "DateTime"
try:
    SQL_WINCC_ROW_LIMIT = int(os.getenv("SQL_WINCC_ROW_LIMIT", "100"))
except ValueError:
    SQL_WINCC_ROW_LIMIT = 100
SQL_WINCC_ROW_LIMIT = max(1, min(SQL_WINCC_ROW_LIMIT, 500))

_VOLTAGE_CHANNEL_COLS_DEFAULT = [
    SQL_VOLTAGE_CH1_COLUMN,
    SQL_VOLTAGE_CH2_COLUMN,
    SQL_VOLTAGE_CH3_COLUMN,
    SQL_VOLTAGE_CH4_COLUMN,
]
_voltage_channel_cols: list[str] = list(_VOLTAGE_CHANNEL_COLS_DEFAULT)
WINDAQ_MSSQL_DEFAULT_EXE = r"C:\WINDAQMSSQL\WINDAQMSSQL.exe"
WINDAQ_MSSQL_EXE = os.getenv("WINDAQ_MSSQL_EXE", "").strip() or WINDAQ_MSSQL_DEFAULT_EXE
WINDAQ_MSSQL_PROCESS_NAME = os.getenv("WINDAQ_MSSQL_PROCESS_NAME", "").strip()
WINDAQ_MSSQL_PID_FILE = os.getenv("WINDAQ_MSSQL_PID_FILE", "").strip() or os.path.join(app.root_path, "windaq_mssql.pid")
WINDAQ_MSSQL_AUTO_START = os.getenv("WINDAQ_MSSQL_AUTO_START", "1").strip().lower() not in {"0", "false", "no"}
WINDAQ_MSSQL_REGISTRY_PATH = r"Software\VB and VBA Program Settings\WINDAQMSSQL\Settings"
WINDAQ_MSSQL_TABLE_SETTING_NAME = "Table Name"

VOLTAGE_TIMESTAMP_COLUMN_CANDIDATES = [
    SQL_VOLTAGE_TIMESTAMP_COLUMN,
    "date_time",
    "DateTime",
    "datetime",
    "timestamp",
    "Timestamp",
    "ReadingTime",
]
VOLTAGE_NON_CHANNEL_COLUMNS = {
    "id",
    "rowid",
    "row_id",
    "recordid",
    "record_id",
    "milliseconds",
    "millisecond",
    "ms",
}
VOLTAGE_NUMERIC_SQL_TYPES = {
    "bigint",
    "decimal",
    "float",
    "int",
    "money",
    "numeric",
    "real",
    "smallint",
    "smallmoney",
    "tinyint",
}

RESULT_ENTRY_FIELDS: list[str] = [
    "time_running",
    "time_flushing",
    "cell_to_cell",
    "bus_2_bus",
    "tank_temp",
    "catholyte_ph",
    "anolyte_ph",
    "catholyte_flow",
    "anolyte_flow",
    "stack_current",
    "current_density",
    "ca_precipitated",
    "mg_precipitated",
    "oh_solution",
    "total_oh",
    "theoretical_oh",
    "catholyte_dosage",
    "anolyte_dosage",
    "ce_oh_titration",
    "cier",
]


def format_timestamp(value: datetime) -> str:
    return value.strftime("%b %d, %Y %I:%M:%S %p")


def export_timestamp(value: datetime) -> str:
    return f"{value.month}/{value.day}/{value.year} {value.strftime('%H:%M:%S')}"


def format_titrator_value(value: Any) -> str:
    if value is None:
        return "—"
    try:
        number = float(value)
    except (TypeError, ValueError):
        return "—"
    text = f"{number:.6f}".rstrip("0").rstrip(".")
    return text or "0"


def bool_status(value: bool | None, true_label: str, false_label: str) -> str:
    if value is None:
        return "—"
    return true_label if value else false_label


def create_seeded_reading(index: int) -> dict[str, Any]:
    timestamp = datetime.now() - timedelta(minutes=(11 - index) * 4)
    r1 = 62 + (index * 0.35) + (7 * math.sin(index / 1.8))
    r2 = 43 + (index * 0.25) + (5 * math.cos(index / 1.6))
    return {
        "timestamp": timestamp,
        "r1": round(r1, 2),
        "r2": round(r2, 2),
        "sequence": "Running" if index > 8 else "Idle",
        "program_a": "Active" if index % 4 == 0 else "Standby",
        "program_b": "Active" if index % 5 == 0 else "Standby",
    }


def update_live_status(live: dict[str, Any]) -> None:
    state.sequence = live["sequence"]
    state.program_a = "Running" if live["prog_a_running"] else "Standby"
    state.program_b = "Running" if live["prog_b_running"] else "Standby"
    state.analyzer_on = live.get("analyzer_on")
    state.analyzer_error = live.get("analyzer_error")
    state.analyzer_failure = live.get("analyzer_failure")
    state.analyzer_timestamp = live.get("analyzer_timestamp")
    state.connection_ok = True


def append_reading() -> dict[str, Any] | None:
    live = read_modbus_live(include_values=not titrator_sql_source_enabled())
    if live:
        update_live_status(live)
        if titrator_sql_source_enabled():
            rows = refresh_titrator_cache()
            return rows[-1] if rows else None
        row = {
            "timestamp": live["analyzer_timestamp"] or datetime.now(),
            "r1": live["r1"],
            "r2": live["r2"],
            "sequence": state.sequence,
            "program_a": state.program_a,
            "program_b": state.program_b,
        }
    elif titrator_sql_source_enabled():
        rows = refresh_titrator_cache()
        if rows:
            state.connection_ok = True
            return rows[-1]
        return None
    else:
        if modbus_enabled():
            state.connection_ok = False
        base_r1 = history[-1]["r1"] if history else 65.0
        base_r2 = history[-1]["r2"] if history else 45.0
        step = 1.4 if state.sequence == "Running" else 0.5
        modifier = random.uniform(-0.8, 0.8)
        row = {
            "timestamp": datetime.now(),
            "r1": round(base_r1 + step + modifier, 2),
            "r2": round(base_r2 + modifier * 1.2 + (0.7 if state.program_b == "Active" else 0.2), 2),
            "sequence": state.sequence,
            "program_a": state.program_a,
            "program_b": state.program_b,
        }
    history.append(row)
    del history[:-12]
    return row


def add_command(command: str, address: str) -> None:
    command_log.insert(
        0,
        {
            "command": command,
            "address": address,
            "time": datetime.now().strftime("%I:%M:%S %p"),
        },
    )
    save_command_log()


def machine_control_actions() -> list[dict[str, str]]:
    tones = {
        "Start Sequence": "start",
        "Stop Sequence": "stop",
        "Break Sequence": "break",
    }
    visible_commands = [
        "Start Sequence",
        "Stop Sequence",
        "Break Sequence",
        "Trigger Program A",
        "Trigger Program B",
    ]
    return [
        {
            "label": command,
            "address": str(COMMAND_COIL_MANUAL_ADDRESS_MAP[command]),
            "tone": tones.get(command, "neutral"),
        }
        for command in visible_commands
    ]


def _connection_text() -> str:
    if modbus_enabled():
        return "Live Modbus data active" if state.connection_ok else "Modbus connection lost"
    if titrator_sql_source_enabled():
        return "SQL Server titrator data active" if state.connection_ok else "SQL Server titrator connection unavailable"
    return "Simulation mode — set MODBUS_HOST or SQL credentials to connect live"


def latest_summary() -> dict[str, Any]:
    analyzer_timestamp = state.analyzer_timestamp
    app_status = titrolyzer_app_status()
    if not history:
        return {
            "connection_text": _connection_text(),
            "titrolyzer_app_running": app_status["running"],
            "titrolyzer_app_status": app_status["label"],
            "titrolyzer_app_detail": app_status["detail"],
            "analyzer_on_status": bool_status(state.analyzer_on, "On", "Off"),
            "analyzer_error_status": bool_status(state.analyzer_error, "Error", "Clear"),
            "analyzer_failure_status": bool_status(state.analyzer_failure, "Failure", "Clear"),
            "sequence_status": state.sequence,
            "program_a_status": state.program_a,
            "program_b_status": state.program_b,
            "analyzer_timestamp": analyzer_timestamp.strftime("%I:%M:%S %p") if analyzer_timestamp else "—",
            "r1_value": "—",
            "r2_value": "—",
            "r1_trend": "No data",
            "r2_trend": "No data",
            "reading_age": "—",
        }

    latest = history[-1]
    age_seconds = max(0, int((datetime.now() - latest["timestamp"]).total_seconds()))
    r1_rows = [row for row in history if row["r1"] is not None]
    r2_rows = [row for row in history if row["r2"] is not None]
    latest_r1 = r1_rows[-1]["r1"] if r1_rows else None
    previous_r1 = r1_rows[-2]["r1"] if len(r1_rows) > 1 else latest_r1
    latest_r2 = r2_rows[-1]["r2"] if r2_rows else None
    previous_r2 = r2_rows[-2]["r2"] if len(r2_rows) > 1 else latest_r2

    def trend(current: float, prior: float) -> str:
        if current > prior:
            return "Upward trend"
        if current < prior:
            return "Slight dip"
        return "Stable"

    return {
        "connection_text": _connection_text(),
        "titrolyzer_app_running": app_status["running"],
        "titrolyzer_app_status": app_status["label"],
        "titrolyzer_app_detail": app_status["detail"],
        "analyzer_on_status": bool_status(state.analyzer_on, "On", "Off"),
        "analyzer_error_status": bool_status(state.analyzer_error, "Error", "Clear"),
        "analyzer_failure_status": bool_status(state.analyzer_failure, "Failure", "Clear"),
        "sequence_status": state.sequence,
        "program_a_status": state.program_a,
        "program_b_status": state.program_b,
        "analyzer_timestamp": (analyzer_timestamp or latest["timestamp"]).strftime("%I:%M:%S %p"),
        "r1_value": format_titrator_value(latest_r1),
        "r2_value": format_titrator_value(latest_r2),
        "r1_trend": trend(latest_r1, previous_r1) if latest_r1 is not None and previous_r1 is not None else "—",
        "r2_trend": trend(latest_r2, previous_r2) if latest_r2 is not None and previous_r2 is not None else "—",
        "reading_age": f"{age_seconds} sec",
    }


def serialize_history(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    return [
        {
            "timestamp": format_timestamp(row["timestamp"]),
            "timestamp_key": titrator_note_key(row["timestamp"]),
            "chart_time": row["timestamp"].strftime("%I:%M %p"),
            "r1": row["r1"],
            "r2": row["r2"],
            "r1_display": format_titrator_value(row["r1"]),
            "r2_display": format_titrator_value(row["r2"]),
            "sequence": row["sequence"],
            "program_a": row["program_a"],
            "program_b": row["program_b"],
            "note": get_titrator_note(row["timestamp"]) or "—",
        }
        for row in rows
    ]


def serialize_chlorine(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    result = []
    for row in rows:
        ts = row["timestamp"]
        if isinstance(ts, str):
            try:
                ts = datetime.fromisoformat(ts)
            except ValueError:
                ts = datetime.now()
        result.append({
            "timestamp": format_timestamp(ts),
            "home_timestamp": ts.strftime("%m/%d/%y %I:%M %p"),
            "timestamp_key": row.get("timestamp_key") or ts.isoformat(timespec="seconds"),
            "value": f'{row["value"]:.2f}',
            "unit": row["unit"],
            "note": row["note"] or "—",
        })
    return result


def _make_result_entry(ts: datetime, **kwargs: str) -> dict[str, Any]:
    return {"timestamp": format_timestamp(ts), "ts_raw": ts, "timestamp_key": ts.isoformat(timespec="seconds"), **kwargs}


def seed_results_entries(now: datetime) -> list[dict[str, Any]]:
    return [
        _make_result_entry(now - timedelta(hours=3),
            time_running="5.25", time_flushing="0.00",
            cell_to_cell="3.59", bus_2_bus="3.60", tank_temp="13.89",
            catholyte_ph="—", anolyte_ph="—",
            catholyte_flow="9994.93", anolyte_flow="7823.11",
            stack_current="1770", current_density="287",
            ca_precipitated="0.00E+00", mg_precipitated="0.00E+00",
            oh_solution="2.23E-15", total_oh="-3.92E-05", theoretical_oh="1.83E-02",
            catholyte_dosage="10.61", anolyte_dosage="8.17",
            ce_oh_titration="91%", cier="1.2336%"),
        _make_result_entry(now - timedelta(hours=2),
            time_running="4.50", time_flushing="0.00",
            cell_to_cell="3.58", bus_2_bus="3.58", tank_temp="12.39",
            catholyte_ph="—", anolyte_ph="—",
            catholyte_flow="10008.12", anolyte_flow="7799.03",
            stack_current="1769", current_density="286",
            ca_precipitated="0.00E+00", mg_precipitated="0.00E+00",
            oh_solution="2.23E-15", total_oh="-3.92E-05", theoretical_oh="1.83E-02",
            catholyte_dosage="10.61", anolyte_dosage="8.16",
            ce_oh_titration="73%", cier="0.0140%"),
        _make_result_entry(now - timedelta(hours=1),
            time_running="4.75", time_flushing="0.00",
            cell_to_cell="3.58", bus_2_bus="3.59", tank_temp="12.93",
            catholyte_ph="—", anolyte_ph="—",
            catholyte_flow="10012.59", anolyte_flow="7801.44",
            stack_current="1768", current_density="286",
            ca_precipitated="0.00E+00", mg_precipitated="0.00E+00",
            oh_solution="2.23E-15", total_oh="-3.92E-05", theoretical_oh="1.83E-02",
            catholyte_dosage="10.59", anolyte_dosage="8.16",
            ce_oh_titration="73%", cier="0.0080%"),
        _make_result_entry(now - timedelta(minutes=15),
            time_running="5.25", time_flushing="0.00",
            cell_to_cell="3.59", bus_2_bus="3.60", tank_temp="13.89",
            catholyte_ph="—", anolyte_ph="—",
            catholyte_flow="9994.93", anolyte_flow="7823.11",
            stack_current="1770", current_density="287",
            ca_precipitated="0.00E+00", mg_precipitated="0.00E+00",
            oh_solution="2.23E-15", total_oh="-3.92E-05", theoretical_oh="1.83E-02",
            catholyte_dosage="10.61", anolyte_dosage="8.17",
            ce_oh_titration="0%", cier="0.0000%"),
    ]


def get_results_rows() -> list[dict[str, str]]:
    return list(reversed(results_entries))


def get_result_cards() -> list[dict[str, Any]]:
    rows = get_results_rows()
    if not rows:
        return []
    latest = rows[0]
    return [
        {"label": "Latest Timestamp", "value": latest["timestamp"], "subtext": "Most recent data row", "tone": "accent-blue", "compact": True},
        {"label": "Current Density", "value": f'{latest["current_density"]} A/m²', "subtext": "Latest reading", "tone": "accent-gold", "compact": False},
        {"label": "CIER", "value": latest["cier"], "subtext": "Latest calculated result", "tone": "accent-green", "compact": False},
        {"label": "CE OH- (Titration)", "value": latest["ce_oh_titration"], "subtext": "Latest calculated result", "tone": "accent-red", "compact": False},
    ]


def simple_xlsx(headers: list[str], rows: list[list[Any]], title: str = "") -> io.BytesIO:
    wb = Workbook()
    ws = wb.active
    if title:
        ws.title = title
    bold = Font(bold=True)
    for col_idx, h in enumerate(headers, start=1):
        c = ws.cell(row=1, column=col_idx, value=h)
        c.font = bold
    for row_idx, row in enumerate(rows, start=2):
        for col_idx, val in enumerate(row, start=1):
            ws.cell(row=row_idx, column=col_idx, value=val)
    for col in ws.columns:
        max_len = max((len(str(c.value or "")) for c in col), default=8)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(max_len + 4, 30)
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output


def csv_bytes(headers: list[str], rows: list[list[Any]]) -> io.BytesIO:
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerow(headers)
    writer.writerows(rows)
    output = io.BytesIO(buffer.getvalue().encode("utf-8"))
    output.seek(0)
    return output


def titrator_note_key(timestamp: datetime) -> str:
    return timestamp.isoformat(timespec="seconds")


def load_titrator_notes() -> None:
    titrator_notes.clear()
    if not os.path.exists(TITRATOR_NOTES_FILE):
        return
    try:
        with open(TITRATOR_NOTES_FILE, "r", encoding="utf-8") as handle:
            payload = json.load(handle)
        if isinstance(payload, dict):
            titrator_notes.update(
                {
                    str(key): str(value).strip()
                    for key, value in payload.items()
                    if str(value).strip()
                }
            )
    except Exception as exc:
        app.logger.warning("Could not load titrator notes: %s", exc)


def save_titrator_notes() -> None:
    try:
        with open(TITRATOR_NOTES_FILE, "w", encoding="utf-8") as handle:
            json.dump(titrator_notes, handle, indent=2)
    except Exception as exc:
        app.logger.warning("Could not save titrator notes: %s", exc)


def load_voltage_notes() -> None:
    voltage_notes.clear()
    if not os.path.exists(VOLTAGE_NOTES_FILE):
        return
    try:
        with open(VOLTAGE_NOTES_FILE, "r", encoding="utf-8") as f:
            payload = json.load(f)
        if isinstance(payload, dict):
            voltage_notes.update({k: v for k, v in payload.items() if str(v).strip()})
    except Exception as exc:
        app.logger.warning("Could not load voltage notes: %s", exc)


def save_voltage_notes() -> None:
    try:
        with open(VOLTAGE_NOTES_FILE, "w", encoding="utf-8") as f:
            json.dump(voltage_notes, f, indent=2)
    except Exception as exc:
        app.logger.warning("Could not save voltage notes: %s", exc)


def load_command_logs() -> None:
    global command_log, voltage_command_log
    for log, path in ((command_log, COMMAND_LOG_FILE), (voltage_command_log, VOLTAGE_COMMAND_LOG_FILE)):
        if not os.path.exists(path):
            continue
        try:
            with open(path, "r", encoding="utf-8") as f:
                data = json.load(f)
            if isinstance(data, list):
                log.extend(data)
        except Exception as exc:
            app.logger.warning("Could not load command log %s: %s", path, exc)


def save_command_log() -> None:
    try:
        with open(COMMAND_LOG_FILE, "w", encoding="utf-8") as f:
            json.dump(command_log, f, indent=2)
    except Exception as exc:
        app.logger.warning("Could not save command log: %s", exc)


def save_voltage_command_log() -> None:
    try:
        with open(VOLTAGE_COMMAND_LOG_FILE, "w", encoding="utf-8") as f:
            json.dump(voltage_command_log, f, indent=2)
    except Exception as exc:
        app.logger.warning("Could not save voltage command log: %s", exc)


def get_titrator_note(timestamp: datetime) -> str:
    return titrator_notes.get(titrator_note_key(timestamp), "")


def set_titrator_note(timestamp: datetime, note: str) -> None:
    key = titrator_note_key(timestamp)
    cleaned = note.strip()
    if cleaned:
        titrator_notes[key] = cleaned
    else:
        titrator_notes.pop(key, None)
    save_titrator_notes()


def parse_titrator_timestamp(value: str) -> datetime | None:
    text = (value or "").strip()
    if not text:
        return None
    try:
        return datetime.fromisoformat(text)
    except ValueError:
        return None


def modbus_enabled() -> bool:
    return bool(MODBUS_HOST and _pymodbus_available)


def _decode_float_le_word(registers: list[int]) -> float:
    import struct
    # Metrohm 2026/2029: Big-Endian bytes, Little-Endian word order.
    # register[0] = low word, register[1] = high word.
    return struct.unpack(">f", struct.pack(">HH", registers[1], registers[0]))[0]


def _decode_uint32_le_word(registers: list[int]) -> int:
    import struct
    return struct.unpack(">I", struct.pack(">HH", registers[1], registers[0]))[0]


def read_modbus_live(include_values: bool = True) -> dict[str, Any] | None:
    if not modbus_enabled():
        return None
    try:
        client = _ModbusTcpClient(MODBUS_HOST, port=MODBUS_PORT)
        if not client.connect():
            return None
        try:
            # Read all status bits: address 3 (manual 4) for 7 bits covers
            # Analyzer On/Error/Failure, Maintenance, Seq Running, Prog A, Prog B
            status_bits = client.read_discrete_inputs(address=3, count=7, device_id=MODBUS_UNIT_ID)
            if status_bits.isError():
                return None
            bits = list(status_bits.bits[:7])
            if len(bits) < 7:
                return None
            seq_running    = bool(bits[4])
            prog_a_running = bool(bits[5])
            prog_b_running = bool(bits[6])

            # Sequence state register: manual address 1 = pymodbus address 0
            ir_seq = client.read_input_registers(address=0, count=1, device_id=MODBUS_UNIT_ID)
            sequence = SEQUENCE_STATE_MAP.get(ir_seq.registers[0], f"State {ir_seq.registers[0]}") if not ir_seq.isError() else ("Running" if seq_running else "Idle")

            r1_value: float | None = None
            r2_value: float | None = None
            if include_values:
                ir_r1 = client.read_input_registers(address=1, count=2, device_id=MODBUS_UNIT_ID)
                if ir_r1.isError():
                    return None

                ir_r2 = client.read_input_registers(address=3, count=2, device_id=MODBUS_UNIT_ID)
                if ir_r2.isError():
                    return None

                r1_value = round(_decode_float_le_word(ir_r1.registers), 2)
                r2_value = round(_decode_float_le_word(ir_r2.registers), 2)

            analyzer_ts: datetime | None = None
            ir_ts = client.read_input_registers(address=37, count=2, device_id=MODBUS_UNIT_ID)
            if not ir_ts.isError():
                epoch = _decode_uint32_le_word(ir_ts.registers)
                if epoch > 0:
                    try:
                        analyzer_ts = datetime.fromtimestamp(epoch, tz=timezone.utc).replace(microsecond=0)
                    except (OSError, OverflowError):
                        pass

            return {
                "analyzer_on": bool(bits[0]),
                "analyzer_error": bool(bits[1]),
                "analyzer_failure": bool(bits[2]),
                "seq_running": seq_running,
                "prog_a_running": prog_a_running,
                "prog_b_running": prog_b_running,
                "sequence": sequence,
                "r1": r1_value,
                "r2": r2_value,
                "analyzer_timestamp": analyzer_ts,
            }
        finally:
            client.close()
    except Exception as exc:
        app.logger.warning("Modbus read failed: %s", exc)
        return None


def write_modbus_coil(command: str) -> bool:
    if not modbus_enabled():
        return False
    coil_address = COMMAND_COIL_MAP.get(command)
    if coil_address is None:
        return False
    try:
        client = _ModbusTcpClient(MODBUS_HOST, port=MODBUS_PORT)
        if not client.connect():
            return False
        try:
            result = client.write_coil(address=coil_address, value=True, device_id=MODBUS_UNIT_ID)
            return not result.isError()
        finally:
            client.close()
    except Exception as exc:
        app.logger.warning("Modbus write coil failed: %s", exc)
        return False


def get_sql_driver() -> str | None:
    if pyodbc is None:
        return None
    drivers = pyodbc.drivers()
    preferred = [name for name in [SQL_DRIVER] if name]
    preferred.extend(
        [
            "ODBC Driver 18 for SQL Server",
            "ODBC Driver 17 for SQL Server",
            "ODBC Driver 13 for SQL Server",
            "SQL Server Native Client 11.0",
            "SQL Server",
        ]
    )
    for name in preferred:
        if name in drivers:
            return name
    return None


def build_sql_connection_string() -> str:
    if SQL_SERVER_CONNECTION_STRING:
        return SQL_SERVER_CONNECTION_STRING
    driver = get_sql_driver()
    if not driver or not SQL_SERVER or not SQL_DATABASE:
        return ""
    parts = [
        f"DRIVER={{{driver}}}",
        f"SERVER={SQL_SERVER}",
        f"DATABASE={SQL_DATABASE}",
    ]
    if SQL_TRUSTED_CONNECTION:
        parts.append("Trusted_Connection=yes")
    elif SQL_USERNAME and SQL_PASSWORD:
        parts.append(f"UID={SQL_USERNAME}")
        parts.append(f"PWD={SQL_PASSWORD}")
    else:
        return ""
    return ";".join(parts) + ";"


def sql_server_enabled() -> bool:
    return bool(build_sql_connection_string())


def titrator_sql_source_enabled() -> bool:
    return sql_server_enabled()


def get_sql_connection() -> Any:
    if not sql_server_enabled():
        return None
    try:
        return pyodbc.connect(build_sql_connection_string(), timeout=5)
    except Exception as exc:
        app.logger.warning("SQL Server connection unavailable: %s", exc)
        return None


def sql_identifier(value: str) -> str:
    return f"[{value.replace(']', ']]')}]"


def sql_qualified_table_name(schema_name: str, table_name: str) -> str:
    return f"{sql_identifier(schema_name)}.{sql_identifier(table_name)}"


def sql_table_name(table_name: str) -> str:
    return sql_qualified_table_name(SQL_SCHEMA, table_name)


def sql_object_name(table_name: str) -> str:
    return f"{SQL_SCHEMA}.{table_name}"


def sql_column_list(columns: list[str]) -> str:
    return ", ".join(sql_identifier(column) for column in columns)


def voltage_note_sql_enabled() -> bool:
    return bool(SQL_VOLTAGE_NOTE_COLUMN)


def table_exists(connection: Any, schema_name: str, table_name: str) -> bool:
    cursor = connection.cursor()
    cursor.execute(
        """
        SELECT 1
        FROM INFORMATION_SCHEMA.TABLES
        WHERE TABLE_SCHEMA = ? AND TABLE_NAME = ?
        """,
        schema_name,
        table_name,
    )
    return cursor.fetchone() is not None


def table_has_columns(connection: Any, schema_name: str, table_name: str, columns: list[str]) -> bool:
    if not columns:
        return True
    cursor = connection.cursor()
    placeholders = ", ".join("?" for _ in columns)
    cursor.execute(
        f"""
        SELECT COLUMN_NAME
        FROM INFORMATION_SCHEMA.COLUMNS
        WHERE TABLE_SCHEMA = ? AND TABLE_NAME = ? AND COLUMN_NAME IN ({placeholders})
        """,
        schema_name,
        table_name,
        *columns,
    )
    found = {row[0] for row in cursor.fetchall()}
    return all(column in found for column in columns)


def split_sql_table_reference(table_ref: str) -> tuple[str, str]:
    parts = [part.strip().strip("[]") for part in table_ref.split(".") if part.strip()]
    if len(parts) >= 2:
        return parts[-2], parts[-1]
    return SQL_SCHEMA, parts[0] if parts else table_ref


def discover_table_columns(connection: Any, schema_name: str, table_name: str) -> list[str]:
    cursor = connection.cursor()
    cursor.execute(
        """
        SELECT c.name
        FROM sys.columns AS c
        INNER JOIN sys.tables AS t ON t.object_id = c.object_id
        INNER JOIN sys.schemas AS s ON s.schema_id = t.schema_id
        WHERE s.name = ? AND t.name = ?
        ORDER BY c.column_id
        """,
        schema_name,
        table_name,
    )
    return [str(row[0]) for row in cursor.fetchall()]


def discover_table_column_types(connection: Any, schema_name: str, table_name: str) -> list[dict[str, str]]:
    cursor = connection.cursor()
    cursor.execute(
        """
        SELECT c.name, ty.name
        FROM sys.columns AS c
        INNER JOIN sys.tables AS t ON t.object_id = c.object_id
        INNER JOIN sys.schemas AS s ON s.schema_id = t.schema_id
        INNER JOIN sys.types AS ty ON ty.user_type_id = c.user_type_id
        WHERE s.name = ? AND t.name = ?
        ORDER BY c.column_id
        """,
        schema_name,
        table_name,
    )
    return [{"name": str(row[0]), "type": str(row[1]).lower()} for row in cursor.fetchall()]


def choose_voltage_timestamp_column(columns: list[str]) -> str | None:
    column_lookup = {column.lower(): column for column in columns}
    for candidate in VOLTAGE_TIMESTAMP_COLUMN_CANDIDATES:
        match = column_lookup.get(candidate.lower())
        if match:
            return match
    for column in columns:
        lowered = column.lower()
        if "date" in lowered or "time" in lowered:
            return column
    return None


def required_voltage_columns() -> list[str]:
    return []


def discover_voltage_channel_columns(connection: Any, table_name: str) -> list[str]:
    column_types = discover_table_column_types(connection, SQL_SCHEMA, table_name)
    table_columns = [column["name"] for column in column_types]
    timestamp_column = choose_voltage_timestamp_column(table_columns)
    excluded = {SQL_VOLTAGE_TIMESTAMP_COLUMN.lower()}
    if timestamp_column:
        excluded.add(timestamp_column.lower())
    for candidate in VOLTAGE_TIMESTAMP_COLUMN_CANDIDATES:
        excluded.add(candidate.lower())
    if SQL_VOLTAGE_ROW_ID_COLUMN:
        excluded.add(SQL_VOLTAGE_ROW_ID_COLUMN.lower())
    if SQL_VOLTAGE_NOTE_COLUMN:
        excluded.add(SQL_VOLTAGE_NOTE_COLUMN.lower())
    excluded.update(VOLTAGE_NON_CHANNEL_COLUMNS)

    channel_columns = [
        column["name"]
        for column in column_types
        if column["name"].lower() not in excluded
        and (
            column["type"] in VOLTAGE_NUMERIC_SQL_TYPES
            or column["name"].lower().startswith(("channel_", "ch"))
        )
    ]
    return channel_columns


def filter_voltage_channels_with_data(connection: Any, table_name: str, channel_cols: list[str]) -> list[str]:
    if not channel_cols:
        return []
    cursor = connection.cursor()
    checks = [
        f"MAX(CASE WHEN {sql_identifier(column)} IS NOT NULL THEN 1 ELSE 0 END)"
        for column in channel_cols
    ]
    cursor.execute(
        f"""
        SELECT COUNT_BIG(1), {", ".join(checks)}
        FROM {sql_table_name(table_name)}
        """
    )
    row = cursor.fetchone()
    if row is None:
        return []
    row_count = int(row[0] or 0)
    if row_count == 0:
        return []
    active = [
        column
        for column, has_data in zip(channel_cols, row[1:])
        if bool(has_data)
    ]
    return active or channel_cols


def voltage_table_row_count(connection: Any, table_name: str) -> int:
    cursor = connection.cursor()
    cursor.execute(f"SELECT COUNT_BIG(1) FROM {sql_table_name(table_name)}")
    row = cursor.fetchone()
    return int(row[0] or 0) if row else 0


def reset_empty_voltage_table_for_windaq(connection: Any) -> bool:
    if not table_exists(connection, SQL_SCHEMA, SQL_VOLTAGE_TABLE):
        return False
    if voltage_table_row_count(connection, SQL_VOLTAGE_TABLE) != 0:
        return False
    cursor = connection.cursor()
    cursor.execute(f"DROP TABLE {sql_table_name(SQL_VOLTAGE_TABLE)}")
    connection.commit()
    return True


_VOLTAGE_BACKUP_TABLE = f"{SQL_VOLTAGE_TABLE}_prev"


def _rename_voltage_to_backup(connection: Any) -> None:
    cursor = connection.cursor()
    if table_exists(connection, SQL_SCHEMA, _VOLTAGE_BACKUP_TABLE):
        cursor.execute(f"DROP TABLE {sql_table_name(_VOLTAGE_BACKUP_TABLE)}")
        connection.commit()
    cursor.execute("EXEC sp_rename ?, ?", (f"{SQL_SCHEMA}.{SQL_VOLTAGE_TABLE}", _VOLTAGE_BACKUP_TABLE))
    connection.commit()


def _merge_voltage_backup(connection: Any) -> None:
    if not table_exists(connection, SQL_SCHEMA, _VOLTAGE_BACKUP_TABLE):
        return
    if not table_exists(connection, SQL_SCHEMA, SQL_VOLTAGE_TABLE):
        cursor = connection.cursor()
        cursor.execute("EXEC sp_rename ?, ?", (f"{SQL_SCHEMA}.{_VOLTAGE_BACKUP_TABLE}", SQL_VOLTAGE_TABLE))
        connection.commit()
        return
    try:
        cursor = connection.cursor()
        cursor.execute("""
            SELECT c1.COLUMN_NAME
            FROM INFORMATION_SCHEMA.COLUMNS c1
            JOIN INFORMATION_SCHEMA.COLUMNS c2
                ON c1.COLUMN_NAME = c2.COLUMN_NAME
                AND c2.TABLE_SCHEMA = ? AND c2.TABLE_NAME = ?
            WHERE c1.TABLE_SCHEMA = ? AND c1.TABLE_NAME = ?
              AND COLUMNPROPERTY(OBJECT_ID(
                    QUOTENAME(c1.TABLE_SCHEMA) + '.' + QUOTENAME(c1.TABLE_NAME)),
                    c1.COLUMN_NAME, 'IsIdentity') = 0
            ORDER BY c1.ORDINAL_POSITION
        """, (SQL_SCHEMA, SQL_VOLTAGE_TABLE, SQL_SCHEMA, _VOLTAGE_BACKUP_TABLE))
        common_cols = [row[0] for row in cursor.fetchall()]
        if common_cols:
            cols_sql = ", ".join(sql_identifier(c) for c in common_cols)
            ts = sql_identifier(SQL_VOLTAGE_TIMESTAMP_COLUMN) if SQL_VOLTAGE_TIMESTAMP_COLUMN in common_cols else None
            order = f"ORDER BY {ts}" if ts else ""
            cursor.execute(f"""
                INSERT INTO {sql_table_name(SQL_VOLTAGE_TABLE)} ({cols_sql})
                SELECT {cols_sql} FROM {sql_table_name(_VOLTAGE_BACKUP_TABLE)} {order}
            """)
            connection.commit()
        cursor.execute(f"DROP TABLE {sql_table_name(_VOLTAGE_BACKUP_TABLE)}")
        connection.commit()
    except Exception as exc:
        app.logger.warning("Could not merge voltage backup: %s", exc)


def _merge_voltage_backup_background() -> None:
    deadline = time.time() + 60
    while time.time() < deadline:
        time.sleep(3)
        connection = get_sql_connection()
        if not connection:
            continue
        try:
            if not table_exists(connection, SQL_SCHEMA, _VOLTAGE_BACKUP_TABLE):
                return
            if table_exists(connection, SQL_SCHEMA, SQL_VOLTAGE_TABLE):
                _merge_voltage_backup(connection)
                return
        except Exception as exc:
            app.logger.warning("Voltage backup merge check failed: %s", exc)
            return
        finally:
            connection.close()


def ensure_voltage_table(
    connection: Any,
    channel_columns: list[str] | None = None,
    create_if_missing: bool = True,
) -> None:
    channel_columns = list(channel_columns or [])
    table_already_exists = table_exists(connection, SQL_SCHEMA, SQL_VOLTAGE_TABLE)
    if not table_already_exists and not create_if_missing:
        return

    column_definitions: list[str] = []
    if SQL_VOLTAGE_ROW_ID_COLUMN:
        column_definitions.append(
            f"{sql_identifier(SQL_VOLTAGE_ROW_ID_COLUMN)} INT IDENTITY(1,1) PRIMARY KEY"
        )
    column_definitions.append(f"{sql_identifier(SQL_VOLTAGE_TIMESTAMP_COLUMN)} DATETIME NOT NULL DEFAULT GETDATE()")
    column_definitions.extend(
        f"{sql_identifier(column)} FLOAT NULL"
        for column in channel_columns
    )
    if voltage_note_sql_enabled():
        column_definitions.append(f"{sql_identifier(SQL_VOLTAGE_NOTE_COLUMN)} NVARCHAR(500) NULL")

    cursor = connection.cursor()
    if not table_already_exists:
        cursor.execute(
            f"""
            CREATE TABLE {sql_table_name(SQL_VOLTAGE_TABLE)} (
                {", ".join(column_definitions)}
            )
            """
        )
        connection.commit()

    existing_column_names = discover_table_columns(connection, SQL_SCHEMA, SQL_VOLTAGE_TABLE)
    existing_columns = {column.lower() for column in existing_column_names}
    expected_columns: list[tuple[str, str]] = []
    if choose_voltage_timestamp_column(existing_column_names) is None:
        expected_columns.append((SQL_VOLTAGE_TIMESTAMP_COLUMN, "DATETIME NULL"))
    expected_columns.extend((column, "FLOAT NULL") for column in channel_columns)
    if voltage_note_sql_enabled():
        expected_columns.append((SQL_VOLTAGE_NOTE_COLUMN, "NVARCHAR(500) NULL"))

    for column, definition in expected_columns:
        if column.lower() in existing_columns:
            continue
        cursor.execute(
            f"""
            ALTER TABLE {sql_table_name(SQL_VOLTAGE_TABLE)}
            ADD {sql_identifier(column)} {definition}
            """
        )
    connection.commit()


def resolve_voltage_table_name(connection: Any) -> str | None:
    required_columns = required_voltage_columns()

    # Prefer the canonical table name when it exists.
    if SQL_VOLTAGE_TABLE and table_exists(connection, SQL_SCHEMA, SQL_VOLTAGE_TABLE):
        ensure_voltage_table(connection, create_if_missing=False)
        if table_has_columns(connection, SQL_SCHEMA, SQL_VOLTAGE_TABLE, required_columns):
            return SQL_VOLTAGE_TABLE

    # Fall back to pattern matching so dated tables (VoltageApr30...) or the
    # backup table (Voltage_prev) are shown whenever the canonical table is absent.
    cursor = connection.cursor()
    cursor.execute(
        """
        SELECT t.name
        FROM sys.tables AS t
        INNER JOIN sys.schemas AS s ON s.schema_id = t.schema_id
        WHERE s.name = ? AND t.name LIKE ?
        ORDER BY t.create_date DESC, t.modify_date DESC, t.name DESC
        """,
        SQL_SCHEMA,
        SQL_VOLTAGE_TABLE_PATTERN,
    )
    candidates = [str(row[0]) for row in cursor.fetchall()]
    # Prefer tables that already contain rows — avoids showing a blank chart
    # while WinDAQ's freshly-created dated table is still empty and the
    # previous data lives in the backup table (Voltage_prev).
    first_empty: str | None = None
    for candidate in candidates:
        if not table_has_columns(connection, SQL_SCHEMA, candidate, required_columns):
            continue
        if voltage_table_row_count(connection, candidate) > 0:
            return candidate
        if first_empty is None:
            first_empty = candidate
    return first_empty


def windows_detached_startup(show_window: bool = False) -> tuple[int, Any]:
    creation_flags = 0
    if hasattr(subprocess, "DETACHED_PROCESS"):
        creation_flags |= subprocess.DETACHED_PROCESS
    if hasattr(subprocess, "CREATE_NEW_PROCESS_GROUP"):
        creation_flags |= subprocess.CREATE_NEW_PROCESS_GROUP
    startupinfo = None
    if os.name == "nt" and hasattr(subprocess, "STARTUPINFO"):
        startupinfo = subprocess.STARTUPINFO()
        startupinfo.dwFlags |= subprocess.STARTF_USESHOWWINDOW
        startupinfo.wShowWindow = 1 if show_window else getattr(subprocess, "SW_HIDE", 0)
    return creation_flags, startupinfo


def read_pid_file(path: str) -> int | None:
    try:
        with open(path, "r", encoding="utf-8") as handle:
            value = int(handle.read().strip())
        return value if value > 0 else None
    except (OSError, ValueError):
        return None


def write_pid_file(path: str, pid: int) -> None:
    with open(path, "w", encoding="utf-8") as handle:
        handle.write(str(pid))


def remove_pid_file(path: str) -> None:
    try:
        os.remove(path)
    except FileNotFoundError:
        pass
    except OSError as exc:
        app.logger.warning("Could not remove PID file %s: %s", path, exc)


def windows_pid_running(pid: int) -> bool:
    if os.name != "nt":
        return False
    try:
        result = subprocess.run(
            ["tasklist", "/FI", f"PID eq {pid}", "/FO", "CSV", "/NH"],
            capture_output=True,
            text=True,
            check=False,
            creationflags=subprocess.CREATE_NO_WINDOW,
        )
    except Exception as exc:
        app.logger.warning("Could not check process %s: %s", pid, exc)
        return False
    output = (result.stdout or "").strip()
    return result.returncode == 0 and str(pid) in output and "INFO:" not in output.upper()


def find_titrolyzer_pid_by_cmdline() -> int | None:
    """Scan running Python processes for a Titrolyzer script by command line using wmic."""
    if os.name != "nt":
        return None
    script_name = os.path.basename(TITROLYZER_APP_PATH) if TITROLYZER_APP_PATH else "Titrolyzer.py"
    try:
        result = subprocess.run(
            [
                "wmic", "process", "where",
                f"(Name='python.exe' or Name='pythonw.exe') and CommandLine like '%{script_name}%'",
                "get", "ProcessId",
            ],
            capture_output=True,
            text=True,
            check=False,
            creationflags=subprocess.CREATE_NO_WINDOW,
        )
    except Exception as exc:
        app.logger.warning("Could not scan processes for Titrolyzer: %s", exc)
        return None
    for line in result.stdout.splitlines():
        line = line.strip()
        if line and line.isdigit():
            pid = int(line)
            if pid > 0:
                return pid
    return None


def titrolyzer_app_status() -> dict[str, Any]:
    if os.name != "nt":
        return {"running": False, "label": "Unavailable", "detail": "Windows only"}
    # Check tracked PID first
    pid = read_pid_file(TITROLYZER_PID_FILE)
    if pid and windows_pid_running(pid):
        return {"running": True, "label": "Running", "detail": f"PID {pid}"}
    # PID file stale or missing — always do a live scan for any Python instance
    if pid:
        remove_pid_file(TITROLYZER_PID_FILE)
    pid = find_titrolyzer_pid_by_cmdline()
    if pid:
        write_pid_file(TITROLYZER_PID_FILE, pid)
        try:
            os.remove(TITROLYZER_STOPPED_FLAG)
        except FileNotFoundError:
            pass
        return {"running": True, "label": "Running", "detail": f"PID {pid} (detected)"}
    return {"running": False, "label": "Closed", "detail": "Not running"}


def start_titrolyzer_app() -> tuple[bool, str]:
    if os.name != "nt":
        return False, "Titrolyzer app start is only supported on Windows."
    if not TITROLYZER_APP_PATH:
        return False, "Set TITROLYZER_APP_PATH so the website knows which Titrolyzer script to launch."
    if not os.path.exists(TITROLYZER_APP_PATH):
        return False, f"Titrolyzer app not found: {TITROLYZER_APP_PATH}"

    existing_pid = read_pid_file(TITROLYZER_PID_FILE)
    if existing_pid and windows_pid_running(existing_pid):
        return True, "Titrolyzer is already running."
    if existing_pid:
        remove_pid_file(TITROLYZER_PID_FILE)
    existing_pid = find_titrolyzer_pid_by_cmdline()
    if existing_pid:
        write_pid_file(TITROLYZER_PID_FILE, existing_pid)
        return True, "Titrolyzer is already running (started externally)."

    if TITROLYZER_APP_PATH.lower().endswith(".py"):
        python_exe = TITROLYZER_PYTHON_EXE or sys.executable
        if not python_exe:
            return False, "Set TITROLYZER_PYTHON_EXE so the website knows which Python to use."
        if os.path.isabs(python_exe) and not os.path.exists(python_exe):
            return False, f"Python executable not found: {python_exe}"
        command = [python_exe, TITROLYZER_APP_PATH]
    else:
        command = [TITROLYZER_APP_PATH]

    try:
        creation_flags, startupinfo = windows_detached_startup()
        process = subprocess.Popen(
            command,
            cwd=os.path.dirname(TITROLYZER_APP_PATH) or None,
            creationflags=creation_flags,
            startupinfo=startupinfo,
            close_fds=True,
        )
        write_pid_file(TITROLYZER_PID_FILE, process.pid)
        try:
            os.remove(TITROLYZER_STOPPED_FLAG)
        except FileNotFoundError:
            pass
        return True, "Recording started."
    except Exception as exc:
        app.logger.warning("Could not start Titrolyzer app: %s", exc)
        return False, f"Could not start Titrolyzer app: {exc}"


def stop_titrolyzer_app() -> tuple[bool, str]:
    if os.name != "nt":
        return False, "Titrolyzer app stop is only supported on Windows."
    pid = read_pid_file(TITROLYZER_PID_FILE)
    if not pid or not windows_pid_running(pid):
        if pid:
            remove_pid_file(TITROLYZER_PID_FILE)
        pid = find_titrolyzer_pid_by_cmdline()
    if not pid:
        try:
            with open(TITROLYZER_STOPPED_FLAG, "w", encoding="utf-8") as f:
                f.write("stopped")
        except OSError:
            pass
        return False, "Titrolyzer is not running."

    try:
        result = subprocess.run(
            ["taskkill", "/PID", str(pid), "/T", "/F"],
            capture_output=True,
            text=True,
            check=False,
            creationflags=subprocess.CREATE_NO_WINDOW,
        )
    except Exception as exc:
        app.logger.warning("Could not stop Titrolyzer app: %s", exc)
        return False, f"Could not stop Titrolyzer app: {exc}"

    if result.returncode == 0:
        remove_pid_file(TITROLYZER_PID_FILE)
        try:
            with open(TITROLYZER_STOPPED_FLAG, "w", encoding="utf-8") as f:
                f.write("stopped")
        except OSError as exc:
            app.logger.warning("Could not write stopped flag: %s", exc)
        return True, "Recording stopped."
    stderr = (result.stderr or result.stdout or "").strip()
    return False, stderr or f"taskkill exited with code {result.returncode}."


def normalize_process_image_name(name: str) -> str:
    return name if name.lower().endswith(".exe") else f"{name}.exe"


def windaq_process_image_names() -> list[str]:
    candidates: list[str] = []

    def add_candidate(name: str) -> None:
        image_name = normalize_process_image_name(name.strip())
        if image_name and image_name.lower() not in {candidate.lower() for candidate in candidates}:
            candidates.append(image_name)

    if WINDAQ_MSSQL_PROCESS_NAME:
        add_candidate(WINDAQ_MSSQL_PROCESS_NAME)
    if WINDAQ_MSSQL_EXE:
        add_candidate(os.path.basename(WINDAQ_MSSQL_EXE))
    add_candidate("WINDAQ-MSSQL.exe")
    add_candidate("WINDAQMSSQL.exe")
    return candidates


def windaq_process_image_name() -> str:
    return windaq_process_image_names()[0]


def find_windaq_mssql_pid() -> int | None:
    if os.name != "nt":
        return None
    for image_name in windaq_process_image_names():
        try:
            result = subprocess.run(
                ["tasklist", "/FI", f"IMAGENAME eq {image_name}", "/FO", "CSV", "/NH"],
                capture_output=True,
                text=True,
                check=False,
                creationflags=subprocess.CREATE_NO_WINDOW,
            )
        except Exception as exc:
            app.logger.warning("Could not scan processes for WINDAQ-MSSQL: %s", exc)
            continue
        if result.returncode != 0 or "INFO:" in (result.stdout or "").upper():
            continue
        for row in csv.reader(io.StringIO(result.stdout or "")):
            if len(row) >= 2 and row[0].lower() == image_name.lower() and row[1].isdigit():
                pid = int(row[1])
                if pid > 0:
                    return pid
    return None


def windaq_mssql_status() -> dict[str, Any]:
    if os.name != "nt":
        return {"running": False, "label": "Unavailable", "detail": "Windows only"}
    pid = read_pid_file(WINDAQ_MSSQL_PID_FILE)
    if pid and windows_pid_running(pid):
        return {"running": True, "label": "Open", "detail": f"PID {pid}"}
    if pid:
        remove_pid_file(WINDAQ_MSSQL_PID_FILE)
    pid = find_windaq_mssql_pid()
    if pid:
        write_pid_file(WINDAQ_MSSQL_PID_FILE, pid)
        return {"running": True, "label": "Open", "detail": f"PID {pid} (detected)"}
    return {"running": False, "label": "Closed", "detail": "Not running"}


def windaq_mssql_table_setting() -> str:
    return f"{SQL_SCHEMA}.{SQL_VOLTAGE_TABLE}" if SQL_SCHEMA else SQL_VOLTAGE_TABLE


def sql_connection_setting_parts() -> dict[str, str]:
    parts: dict[str, str] = {}
    if SQL_SERVER_CONNECTION_STRING:
        for chunk in SQL_SERVER_CONNECTION_STRING.split(";"):
            if "=" not in chunk:
                continue
            key, value = chunk.split("=", 1)
            normalized_key = key.strip().lower().replace(" ", "")
            parts[normalized_key] = value.strip()

    return {
        "server": SQL_SERVER or parts.get("server") or parts.get("datasource") or "",
        "database": SQL_DATABASE or parts.get("database") or parts.get("initialcatalog") or "",
        "username": SQL_USERNAME or parts.get("uid") or parts.get("userid") or "",
        "password": SQL_PASSWORD or parts.get("pwd") or parts.get("password") or "",
    }


def configure_windaq_mssql_table_name() -> tuple[bool, str]:
    if os.name != "nt":
        return False, "WINDAQ-MSSQL table setup is only supported on Windows."
    try:
        import winreg
    except ImportError:
        return False, "The website could not access Windows settings for WINDAQ-MSSQL."

    table_setting = windaq_mssql_table_setting()
    sql_parts = sql_connection_setting_parts()
    try:
        with winreg.CreateKey(winreg.HKEY_CURRENT_USER, WINDAQ_MSSQL_REGISTRY_PATH) as key:
            if sql_parts["server"]:
                winreg.SetValueEx(key, "Server Name", 0, winreg.REG_SZ, sql_parts["server"])
            if sql_parts["username"]:
                winreg.SetValueEx(key, "User Name", 0, winreg.REG_SZ, sql_parts["username"])
            if sql_parts["database"]:
                winreg.SetValueEx(key, "Database", 0, winreg.REG_SZ, sql_parts["database"])
            winreg.SetValueEx(
                key,
                WINDAQ_MSSQL_TABLE_SETTING_NAME,
                0,
                winreg.REG_SZ,
                table_setting,
            )
        return True, table_setting
    except Exception as exc:
        app.logger.warning("Could not set WINDAQ-MSSQL table setting: %s", exc)
        return False, (
            f"The website could not set WINDAQ-MSSQL's Table Name setting. Open "
            f"WINDAQ-MSSQL manually and set Table Name to {table_setting}."
        )


def _win32_text(hwnd: int) -> str:
    import ctypes

    user32 = ctypes.windll.user32
    length = user32.GetWindowTextLengthW(hwnd)
    buffer = ctypes.create_unicode_buffer(length + 1)
    user32.GetWindowTextW(hwnd, buffer, len(buffer))
    return buffer.value


def _win32_class_name(hwnd: int) -> str:
    import ctypes

    buffer = ctypes.create_unicode_buffer(256)
    ctypes.windll.user32.GetClassNameW(hwnd, buffer, len(buffer))
    return buffer.value


def _win32_enum_windows() -> list[int]:
    import ctypes

    user32 = ctypes.windll.user32
    handles: list[int] = []
    callback_type = ctypes.WINFUNCTYPE(ctypes.c_bool, ctypes.c_void_p, ctypes.c_void_p)

    def _callback(hwnd: int, _lparam: int) -> bool:
        handles.append(int(hwnd))
        return True

    callback = callback_type(_callback)
    user32.EnumWindows(callback, 0)
    return handles


def _win32_enum_child_windows(hwnd: int) -> list[int]:
    import ctypes

    user32 = ctypes.windll.user32
    handles: list[int] = []
    callback_type = ctypes.WINFUNCTYPE(ctypes.c_bool, ctypes.c_void_p, ctypes.c_void_p)

    def _callback(child_hwnd: int, _lparam: int) -> bool:
        handles.append(int(child_hwnd))
        return True

    callback = callback_type(_callback)
    user32.EnumChildWindows(hwnd, callback, 0)
    return handles


def _win32_window_pid(hwnd: int) -> int:
    import ctypes

    process_id = ctypes.c_ulong()
    ctypes.windll.user32.GetWindowThreadProcessId(hwnd, ctypes.byref(process_id))
    return int(process_id.value)


def _win32_window_rect(hwnd: int) -> tuple[int, int, int, int]:
    import ctypes
    from ctypes import wintypes

    rect = wintypes.RECT()
    ctypes.windll.user32.GetWindowRect(hwnd, ctypes.byref(rect))
    return int(rect.left), int(rect.top), int(rect.right), int(rect.bottom)


def _win32_find_windaq_window(pid: int | None = None, timeout_seconds: float = 8.0) -> int | None:
    import ctypes

    user32 = ctypes.windll.user32
    deadline = time.time() + timeout_seconds
    while time.time() < deadline:
        for hwnd in _win32_enum_windows():
            if not user32.IsWindowVisible(hwnd):
                continue
            title = _win32_text(hwnd).strip()
            if "windaq-mssql" not in title.lower():
                continue
            if pid and _win32_window_pid(hwnd) != pid:
                continue
            return hwnd
        time.sleep(0.25)
    return None


def _win32_find_child_button(hwnd: int, caption: str, require_enabled: bool = True) -> int | None:
    import ctypes

    user32 = ctypes.windll.user32
    for child in _win32_enum_child_windows(hwnd):
        child_text = _win32_text(child).strip().replace("&", "")
        if child_text.lower() != caption.lower():
            continue
        class_name = _win32_class_name(child).lower()
        if "button" not in class_name and "command" not in class_name and "check" not in class_name:
            continue
        if not user32.IsWindowVisible(child):
            continue
        if require_enabled and not user32.IsWindowEnabled(child):
            continue
        return child
    return None


def _win32_click(hwnd: int) -> None:
    import ctypes

    ctypes.windll.user32.SendMessageW(hwnd, 0x00F5, 0, 0)  # BM_CLICK


def _win32_select_first_list_item(hwnd: int) -> bool:
    import ctypes

    user32 = ctypes.windll.user32
    for child in _win32_enum_child_windows(hwnd):
        class_name = _win32_class_name(child).lower()
        if "listbox" not in class_name:
            continue
        count = user32.SendMessageW(child, 0x018B, 0, 0)  # LB_GETCOUNT
        if count <= 0:
            return False
        user32.SendMessageW(child, 0x0186, 0, 0)  # LB_SETCURSEL
        return True
    return False


def _win32_set_windaq_text_fields(hwnd: int) -> list[str]:
    import ctypes

    user32 = ctypes.windll.user32
    sql_parts = sql_connection_setting_parts()
    table_setting = windaq_mssql_table_setting()
    warnings: list[str] = []
    edits: list[dict[str, Any]] = []
    for child in _win32_enum_child_windows(hwnd):
        class_name = _win32_class_name(child).lower()
        if "textbox" not in class_name and "edit" not in class_name:
            continue
        left, top, right, bottom = _win32_window_rect(child)
        edits.append({"hwnd": child, "left": left, "top": top, "right": right, "bottom": bottom})

    if not edits:
        return ["could not find the WINDAQ-MSSQL text fields"]

    main_left, _, main_right, _ = _win32_window_rect(hwnd)
    midpoint = main_left + ((main_right - main_left) / 2)
    left_edits = sorted([edit for edit in edits if edit["left"] < midpoint], key=lambda edit: edit["top"])
    right_edits = sorted([edit for edit in edits if edit["left"] >= midpoint], key=lambda edit: edit["top"])

    def set_field(field_hwnd: int, value: str) -> None:
        # EM_REPLACESEL fires EN_CHANGE so VB6 updates its Text property.
        # WM_SETTEXT (SetWindowTextW) does not fire EN_CHANGE and VB6 ignores it.
        buf = ctypes.create_unicode_buffer(value)
        user32.SendMessageW(field_hwnd, 0x00B1, 0, -1)  # EM_SETSEL: select all
        user32.SendMessageW(field_hwnd, 0x00C2, 0, buf)  # EM_REPLACESEL

    if len(left_edits) >= 4:
        field_values = [
            sql_parts["server"],
            sql_parts["username"],
            sql_parts["password"],
            sql_parts["database"],
        ]
        for edit, value in zip(left_edits[:4], field_values):
            if value:
                set_field(edit["hwnd"], value)
        if not sql_parts["password"]:
            warnings.append("SQL password was not available to fill into WINDAQ-MSSQL")
    else:
        warnings.append("could not fill all WINDAQ-MSSQL SQL connection fields")

    if right_edits:
        set_field(right_edits[-1]["hwnd"], table_setting)
        if len(right_edits) >= 2:
            set_field(right_edits[0]["hwnd"], "1000")
    else:
        warnings.append("could not fill the WINDAQ-MSSQL table name")

    return warnings


def auto_start_windaq_mssql_window(pid: int | None) -> tuple[bool, str]:
    if os.name != "nt":
        return False, "WINDAQ-MSSQL auto-start is only supported on Windows."
    if not WINDAQ_MSSQL_AUTO_START:
        return True, "WINDAQ-MSSQL opened. Auto-start is turned off."

    try:
        import ctypes

        user32 = ctypes.windll.user32
        hwnd = _win32_find_windaq_window(pid)
        if not hwnd:
            return False, "WINDAQ-MSSQL opened, but the website could not find its window to press Connect and Start."

        user32.ShowWindow(hwnd, 9)  # SW_RESTORE
        user32.SetForegroundWindow(hwnd)

        stop_button = _win32_find_child_button(hwnd, "Stop", require_enabled=False)
        if stop_button and user32.IsWindowEnabled(stop_button):
            return True, "WINDAQ-MSSQL already appears to be recording."

        warnings = _win32_set_windaq_text_fields(hwnd)
        selected_device = _win32_select_first_list_item(hwnd)

        # Find the Enable Looping checkbox by class name — VB6 ThunderRT6CheckBox
        # returns empty text from GetWindowTextW, so caption matching doesn't work.
        enable_looping = next(
            (
                child for child in _win32_enum_child_windows(hwnd)
                if "check" in _win32_class_name(child).lower()
                and user32.IsWindowVisible(child)
                and user32.IsWindowEnabled(child)
            ),
            None,
        )
        if enable_looping:
            if user32.SendMessageW(enable_looping, 0x00F0, 0, 0) == 0:  # BM_GETCHECK == unchecked
                _win32_click(enable_looping)
                time.sleep(0.3)
        else:
            app.logger.warning("auto_start_windaq: Enable Looping checkbox not found")

        connect_button = _win32_find_child_button(hwnd, "Connect")
        if connect_button:
            _win32_click(connect_button)
            time.sleep(1.5)

        deadline = time.time() + 6.0
        start_button = _win32_find_child_button(hwnd, "Start")
        while not start_button and time.time() < deadline:
            if stop_button and user32.IsWindowEnabled(stop_button):
                return True, "WINDAQ-MSSQL already appears to be recording."
            time.sleep(0.25)
            start_button = _win32_find_child_button(hwnd, "Start")

        if not start_button:
            detail = " ".join(warnings)
            if not selected_device:
                detail = (detail + " No hardware device was available to select.").strip()
            return False, (
                "WINDAQ-MSSQL opened, but the website could not press Start. "
                "Check the SQL login fields and make sure a DATAQ device is listed. "
                + detail
            ).strip()

        _win32_click(start_button)
        if warnings:
            return True, "WinDAQ-MSSQL started. Note: " + " ".join(warnings)
        return True, "WinDAQ-MSSQL started."
    except Exception as exc:
        app.logger.warning("Could not auto-start WINDAQ-MSSQL window: %s", exc)
        return False, f"WINDAQ-MSSQL opened, but the website could not press Connect/Start automatically: {exc}"


def prepare_voltage_table_for_windaq_start() -> tuple[bool, str]:
    if not sql_server_enabled():
        return False, f"SQL Server is not configured, so the website could not prepare {SQL_SCHEMA}.{SQL_VOLTAGE_TABLE} before starting WINDAQ-MSSQL."
    connection = get_sql_connection()
    if connection is None:
        return False, f"SQL Server connection unavailable, so WINDAQ-MSSQL was not started. The website needs to prepare {SQL_SCHEMA}.{SQL_VOLTAGE_TABLE} first."
    try:
        if table_exists(connection, SQL_SCHEMA, SQL_VOLTAGE_TABLE):
            if voltage_table_row_count(connection, SQL_VOLTAGE_TABLE) == 0:
                reset_empty_voltage_table_for_windaq(connection)
            else:
                _rename_voltage_to_backup(connection)
        configured, message = configure_windaq_mssql_table_name()
        if not configured:
            return False, message
        return True, ""
    except Exception as exc:
        app.logger.warning("Could not prepare voltage SQL table: %s", exc)
        return False, f"Could not prepare {SQL_SCHEMA}.{SQL_VOLTAGE_TABLE} before starting WINDAQ-MSSQL."
    finally:
        connection.close()


def start_windaq_mssql() -> tuple[bool, str]:
    if os.name != "nt":
        return False, "WINDAQ-MSSQL start is only supported on Windows."

    existing_pid = read_pid_file(WINDAQ_MSSQL_PID_FILE)
    if existing_pid and windows_pid_running(existing_pid):
        configure_windaq_mssql_table_name()
        started, message = auto_start_windaq_mssql_window(existing_pid)
        return started, message
    if existing_pid:
        remove_pid_file(WINDAQ_MSSQL_PID_FILE)
    existing_pid = find_windaq_mssql_pid()
    if existing_pid:
        write_pid_file(WINDAQ_MSSQL_PID_FILE, existing_pid)
        configure_windaq_mssql_table_name()
        started, message = auto_start_windaq_mssql_window(existing_pid)
        return started, message

    if not os.path.exists(WINDAQ_MSSQL_EXE):
        return False, f"WINDAQ-MSSQL is not running, and the launch shortcut points to a file that was not found: {WINDAQ_MSSQL_EXE}. Start it manually, or update WINDAQ_MSSQL_EXE to the correct install path."

    prepared, prepare_message = prepare_voltage_table_for_windaq_start()
    if not prepared:
        return False, prepare_message

    try:
        creation_flags, startupinfo = windows_detached_startup(show_window=True)
        process = subprocess.Popen(
            [WINDAQ_MSSQL_EXE],
            cwd=os.path.dirname(WINDAQ_MSSQL_EXE) or None,
            creationflags=creation_flags,
            startupinfo=startupinfo,
            close_fds=True,
        )
        write_pid_file(WINDAQ_MSSQL_PID_FILE, process.pid)
        started, message = auto_start_windaq_mssql_window(process.pid)
        if started:
            threading.Thread(target=_merge_voltage_backup_background, daemon=True).start()
        return started, message
    except Exception as exc:
        app.logger.warning("Could not start WINDAQ-MSSQL: %s", exc)
        return False, f"Could not start WINDAQ-MSSQL: {exc}"


def stop_windaq_mssql() -> tuple[bool, str]:
    if os.name != "nt":
        return False, "WINDAQ-MSSQL stop is only supported on Windows."
    pid = read_pid_file(WINDAQ_MSSQL_PID_FILE)
    if not pid or not windows_pid_running(pid):
        if pid:
            remove_pid_file(WINDAQ_MSSQL_PID_FILE)
        pid = find_windaq_mssql_pid()
    if not pid:
        return False, "WINDAQ-MSSQL is not running."

    try:
        result = subprocess.run(
            ["taskkill", "/PID", str(pid), "/T", "/F"],
            capture_output=True,
            text=True,
            check=False,
            creationflags=subprocess.CREATE_NO_WINDOW,
        )
    except Exception as exc:
        app.logger.warning("Could not stop WINDAQ-MSSQL: %s", exc)
        return False, f"Could not stop WINDAQ-MSSQL: {exc}"

    if result.returncode == 0:
        remove_pid_file(WINDAQ_MSSQL_PID_FILE)
        return True, "WINDAQ-MSSQL stop requested."
    stderr = (result.stderr or result.stdout or "").strip()
    if "not found" in stderr.lower():
        remove_pid_file(WINDAQ_MSSQL_PID_FILE)
        return False, "WINDAQ-MSSQL is not running."
    return False, stderr or f"taskkill exited with code {result.returncode}."


def text_or_default(value: Any, default: str = "—") -> str:
    if value is None:
        return default
    text = str(value).strip()
    return text or default


def _safe_float(value: Any) -> float:
    try:
        return float(value)
    except (TypeError, ValueError):
        return 0.0


def coerce_datetime(value: Any) -> datetime | None:
    if isinstance(value, datetime):
        return value
    if isinstance(value, str):
        try:
            return datetime.fromisoformat(value)
        except ValueError:
            return None
    return None


def format_wincc_column_label(column_name: str) -> str:
    return column_name.replace("_", " ")


def format_wincc_value(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%m/%d/%y %H:%M:%S")
    if isinstance(value, float):
        text = f"{value:.6f}".rstrip("0").rstrip(".")
        return text or "0"
    return text_or_default(value, "")


def format_wincc_home_timestamp(value: Any) -> str:
    if value is None:
        return "No rows"
    if isinstance(value, datetime):
        return value.strftime("%m/%d/%y %H:%M")
    return format_wincc_value(value) or "No rows"


def empty_wincc_source(table_ref: str, error: str = "") -> dict[str, Any]:
    schema_name, table_name = split_sql_table_reference(table_ref)
    return {
        "table_ref": table_ref,
        "schema": schema_name,
        "table": table_name,
        "qualified_name": f"{schema_name}.{table_name}",
        "columns": [],
        "rows": [],
        "latest": [],
        "row_count": 0,
        "home_status": "Unavailable" if error else "No rows",
        "error": error,
    }


def wincc_error_sources(message: str) -> list[dict[str, Any]]:
    return [empty_wincc_source(table_ref, message) for table_ref in SQL_WINCC_TABLE_NAMES]


def load_wincc_table_source(connection: Any, table_ref: str) -> dict[str, Any]:
    schema_name, table_name = split_sql_table_reference(table_ref)
    source = empty_wincc_source(table_ref)

    if not table_exists(connection, schema_name, table_name):
        source["error"] = f"{source['qualified_name']} was not found."
        return source

    columns = discover_table_columns(connection, schema_name, table_name)
    if not columns:
        source["error"] = f"{source['qualified_name']} has no columns."
        return source

    timestamp_column = next(
        (column for column in columns if column.lower() == SQL_WINCC_TIMESTAMP_COLUMN.lower()),
        columns[0],
    )
    timestamp_index = columns.index(timestamp_column)
    cursor = connection.cursor()
    cursor.execute(
        f"""
        SELECT TOP ({SQL_WINCC_ROW_LIMIT})
            {sql_column_list(columns)}
        FROM {sql_qualified_table_name(schema_name, table_name)}
        ORDER BY {sql_identifier(timestamp_column)} DESC
        """
    )

    db_rows = cursor.fetchall()
    rows = [
        [format_wincc_value(value) for value in db_row]
        for db_row in db_rows
    ]
    latest_timestamp_value = db_rows[0][timestamp_index] if db_rows else None
    column_meta = [{"name": column, "label": format_wincc_column_label(column)} for column in columns]
    latest = [
        {"label": meta["label"], "value": value}
        for meta, value in zip(column_meta, rows[0])
    ][:8] if rows else []

    source.update({
        "columns": column_meta,
        "rows": rows,
        "latest": latest,
        "latest_timestamp": format_wincc_value(latest_timestamp_value),
        "home_status": format_wincc_home_timestamp(latest_timestamp_value),
        "row_count": len(rows),
    })
    return source


def load_wincc_sources() -> list[dict[str, Any]]:
    if not sql_server_enabled():
        return wincc_error_sources("SQL Server is not configured.")

    connection = get_sql_connection()
    if connection is None:
        return wincc_error_sources("SQL Server connection unavailable.")

    try:
        sources = []
        for table_ref in SQL_WINCC_TABLE_NAMES:
            try:
                sources.append(load_wincc_table_source(connection, table_ref))
            except Exception as exc:
                app.logger.warning("Could not load WinCC table %s: %s", table_ref, exc)
                sources.append(empty_wincc_source(table_ref, "Could not load this table."))
        return sources
    finally:
        connection.close()


def clear_wincc_table(table_ref: str) -> tuple[bool, str]:
    if table_ref not in SQL_WINCC_TABLE_NAMES:
        return False, "Unknown WinCC table."

    if not sql_server_enabled():
        return False, "SQL Server is not configured."

    connection = get_sql_connection()
    if connection is None:
        return False, "SQL Server connection unavailable."

    schema_name, table_name = split_sql_table_reference(table_ref)
    qualified_name = f"{schema_name}.{table_name}"
    try:
        if not table_exists(connection, schema_name, table_name):
            return False, f"{qualified_name} was not found."
        cursor = connection.cursor()
        cursor.execute(f"DELETE FROM {sql_qualified_table_name(schema_name, table_name)}")
        connection.commit()
        return True, f"{qualified_name} cleared."
    except Exception as exc:
        app.logger.warning("Could not clear WinCC table %s: %s", table_ref, exc)
        return False, f"{qualified_name} could not be cleared."
    finally:
        connection.close()


def safe_worksheet_title(title: str, used_titles: set[str]) -> str:
    cleaned = "".join("_" if char in "[]:*?/\\" else char for char in title).strip() or "Sheet"
    cleaned = cleaned[:31]
    candidate = cleaned
    index = 2
    while candidate in used_titles:
        suffix = f" {index}"
        candidate = f"{cleaned[:31 - len(suffix)]}{suffix}"
        index += 1
    used_titles.add(candidate)
    return candidate


def ensure_chlorine_table() -> None:
    connection = get_sql_connection()
    if connection is None:
        return

    try:
        cursor = connection.cursor()
        cursor.execute(
            f"""
            IF OBJECT_ID('{sql_object_name(SQL_CHLORINE_TABLE)}', 'U') IS NULL
            BEGIN
                CREATE TABLE {sql_table_name(SQL_CHLORINE_TABLE)} (
                    id INT IDENTITY(1,1) PRIMARY KEY,
                    measured_at DATETIME2 NOT NULL,
                    chlorine_value DECIMAL(10, 2) NOT NULL,
                    unit NVARCHAR(20) NOT NULL CONSTRAINT DF_Chlorine_unit DEFAULT 'mg/L',
                    note NVARCHAR(1000) NULL,
                    created_at DATETIME2 NOT NULL CONSTRAINT DF_Chlorine_created_at DEFAULT SYSUTCDATETIME()
                )
            END
            """
        )
        connection.commit()
    finally:
        connection.close()


def load_chlorine_measurements() -> list[dict[str, Any]]:
    fallback = sorted(chlorine_measurements, key=lambda row: row["timestamp"], reverse=True)
    if not sql_server_enabled():
        return fallback

    ensure_chlorine_table()
    connection = get_sql_connection()
    if connection is None:
        return fallback

    try:
        cursor = connection.cursor()
        cursor.execute(
            f"""
            SELECT TOP (100)
                measured_at,
                chlorine_value,
                unit,
                note
            FROM {sql_table_name(SQL_CHLORINE_TABLE)}
            ORDER BY measured_at DESC, id DESC
            """
        )
        rows = []
        for measured_at, chlorine_value, unit, note in cursor.fetchall():
            ts = coerce_datetime(measured_at) or datetime.now()
            rows.append(
                {
                    "timestamp": ts,
                    "timestamp_key": ts.isoformat(timespec="seconds"),
                    "value": float(chlorine_value),
                    "unit": unit,
                    "note": note or "",
                }
            )
        return rows
    except Exception as exc:
        app.logger.warning("Could not load chlorine measurements from SQL Server: %s", exc)
        return fallback
    finally:
        connection.close()


def save_chlorine_measurement(timestamp: datetime, value: float, unit: str, note: str) -> None:
    if not sql_server_enabled():
        chlorine_measurements.append(
            {
                "timestamp": timestamp,
                "timestamp_key": timestamp.isoformat(timespec="seconds"),
                "value": value,
                "unit": unit,
                "note": note,
            }
        )
        chlorine_measurements.sort(key=lambda row: row["timestamp"], reverse=True)
        return

    try:
        ensure_chlorine_table()
        connection = get_sql_connection()
        if connection is None:
            raise RuntimeError("SQL connection unavailable")
        cursor = connection.cursor()
        cursor.execute(
            f"""
            INSERT INTO {sql_table_name(SQL_CHLORINE_TABLE)} (measured_at, chlorine_value, unit, note)
            VALUES (?, ?, ?, ?)
            """,
            timestamp,
            value,
            unit,
            note or None,
        )
        connection.commit()
    except Exception as exc:
        app.logger.warning("Could not save chlorine measurement to SQL Server: %s", exc)
        chlorine_measurements.append(
            {
                "timestamp": timestamp,
                "value": value,
                "unit": unit,
                "note": note,
            }
        )
        chlorine_measurements.sort(key=lambda row: row["timestamp"], reverse=True)
    finally:
        if "connection" in locals() and connection is not None:
            connection.close()


def clear_chlorine_storage() -> bool:
    if not sql_server_enabled():
        chlorine_measurements.clear()
        return True

    ensure_chlorine_table()
    connection = get_sql_connection()
    if connection is None:
        state.connection_ok = False
        return False

    try:
        cursor = connection.cursor()
        cursor.execute(f"DELETE FROM {sql_table_name(SQL_CHLORINE_TABLE)}")
        connection.commit()
        chlorine_measurements.clear()
        return True
    except Exception as exc:
        app.logger.warning("Could not clear chlorine SQL table: %s", exc)
        return False
    finally:
        connection.close()


def ensure_results_table() -> None:
    connection = get_sql_connection()
    if connection is None:
        return

    try:
        cursor = connection.cursor()
        cursor.execute(
            f"""
            IF OBJECT_ID('{sql_object_name(SQL_RESULTS_TABLE)}', 'U') IS NULL
            BEGIN
                CREATE TABLE {sql_table_name(SQL_RESULTS_TABLE)} (
                    id INT IDENTITY(1,1) PRIMARY KEY,
                    measured_at DATETIME2 NOT NULL,
                    time_running NVARCHAR(50) NULL,
                    time_flushing NVARCHAR(50) NULL,
                    cell_to_cell NVARCHAR(50) NULL,
                    bus_2_bus NVARCHAR(50) NULL,
                    tank_temp NVARCHAR(50) NULL,
                    catholyte_ph NVARCHAR(50) NULL,
                    anolyte_ph NVARCHAR(50) NULL,
                    catholyte_flow NVARCHAR(50) NULL,
                    anolyte_flow NVARCHAR(50) NULL,
                    stack_current NVARCHAR(50) NULL,
                    current_density NVARCHAR(50) NULL,
                    ca_precipitated NVARCHAR(50) NULL,
                    mg_precipitated NVARCHAR(50) NULL,
                    oh_solution NVARCHAR(50) NULL,
                    total_oh NVARCHAR(50) NULL,
                    theoretical_oh NVARCHAR(50) NULL,
                    catholyte_dosage NVARCHAR(50) NULL,
                    anolyte_dosage NVARCHAR(50) NULL,
                    ce_oh_titration NVARCHAR(50) NULL,
                    cier NVARCHAR(50) NULL,
                    note NVARCHAR(1000) NULL,
                    created_at DATETIME2 NOT NULL CONSTRAINT DF_Results_created_at DEFAULT SYSUTCDATETIME()
                )
            END
            """
        )
        connection.commit()
    finally:
        connection.close()


def load_results_entries() -> list[dict[str, Any]]:
    fallback = list(results_entries)
    if not sql_server_enabled():
        return fallback

    ensure_results_table()
    connection = get_sql_connection()
    if connection is None:
        return fallback

    try:
        cursor = connection.cursor()
        cursor.execute(
            f"""
            SELECT TOP (200)
                measured_at,
                {sql_column_list(RESULT_ENTRY_FIELDS)},
                note
            FROM {sql_table_name(SQL_RESULTS_TABLE)}
            ORDER BY measured_at DESC, id DESC
            """
        )
        rows: list[dict[str, Any]] = []
        for db_row in cursor.fetchall():
            measured_at = db_row[0]
            entry = {"timestamp": format_timestamp(measured_at), "ts_raw": measured_at, "timestamp_key": measured_at.isoformat(timespec="seconds")}
            for index, field in enumerate(RESULT_ENTRY_FIELDS, start=1):
                entry[field] = text_or_default(db_row[index])
            entry["note"] = text_or_default(db_row[len(RESULT_ENTRY_FIELDS) + 1], "")
            rows.append(entry)
        return list(reversed(rows))
    except Exception as exc:
        app.logger.warning("Could not load results from SQL Server: %s", exc)
        return fallback
    finally:
        connection.close()


def save_results_entry(entry: dict[str, Any]) -> None:
    if not sql_server_enabled():
        results_entries.append(entry)
        return

    columns = ["measured_at", *RESULT_ENTRY_FIELDS, "note"]
    values = [entry["ts_raw"], *(entry.get(field) for field in RESULT_ENTRY_FIELDS), entry.get("note") or None]
    placeholders = ", ".join("?" for _ in columns)
    try:
        ensure_results_table()
        connection = get_sql_connection()
        if connection is None:
            raise RuntimeError("SQL connection unavailable")
        cursor = connection.cursor()
        cursor.execute(
            f"""
            INSERT INTO {sql_table_name(SQL_RESULTS_TABLE)} ({sql_column_list(columns)})
            VALUES ({placeholders})
            """,
            *values,
        )
        connection.commit()
    except Exception as exc:
        app.logger.warning("Could not save results entry to SQL Server: %s", exc)
        results_entries.append(entry)
    finally:
        if "connection" in locals() and connection is not None:
            connection.close()


def clear_results_storage() -> bool:
    if not sql_server_enabled():
        results_entries.clear()
        return True

    ensure_results_table()
    connection = get_sql_connection()
    if connection is None:
        return False

    try:
        cursor = connection.cursor()
        cursor.execute(f"DELETE FROM {sql_table_name(SQL_RESULTS_TABLE)}")
        connection.commit()
        results_entries.clear()
        return True
    except Exception as exc:
        app.logger.warning("Could not clear results SQL table: %s", exc)
        return False
    finally:
        connection.close()


def ensure_titrator_table(connection: Any) -> None:
    cursor = connection.cursor()
    cursor.execute(
        f"""
        IF OBJECT_ID('{sql_object_name(SQL_TITRATOR_TABLE)}', 'U') IS NULL
        BEGIN
            CREATE TABLE {sql_table_name(SQL_TITRATOR_TABLE)} (
                {sql_identifier(SQL_TITRATOR_RECORD_ID_COLUMN)} INT IDENTITY(1,1) PRIMARY KEY,
                {sql_identifier(SQL_TITRATOR_READING_TIME_COLUMN)} DATETIME NOT NULL DEFAULT GETDATE(),
                {sql_identifier(SQL_TITRATOR_TIMESTAMP_COLUMN)} DATETIME NULL,
                {sql_identifier(SQL_TITRATOR_RESULT_ID_COLUMN)} NVARCHAR(5) NULL,
                {sql_identifier(SQL_TITRATOR_VALUE_COLUMN)} FLOAT NULL
            )
        END
        """
    )
    connection.commit()


def load_titrator_history() -> list[dict[str, Any]]:
    fallback = list(history)
    if not titrator_sql_source_enabled():
        return fallback

    connection = get_sql_connection()
    if connection is None:
        state.connection_ok = False
        return fallback

    try:
        ensure_titrator_table(connection)
        cursor = connection.cursor()
        cursor.execute(
            f"""
            SELECT TOP (200)
                {sql_column_list([
                    SQL_TITRATOR_READING_TIME_COLUMN,
                    SQL_TITRATOR_TIMESTAMP_COLUMN,
                    SQL_TITRATOR_RESULT_ID_COLUMN,
                    SQL_TITRATOR_VALUE_COLUMN,
                ])}
            FROM {sql_table_name(SQL_TITRATOR_TABLE)}
            ORDER BY {sql_identifier(SQL_TITRATOR_READING_TIME_COLUMN)} DESC,
                     {sql_identifier(SQL_TITRATOR_RECORD_ID_COLUMN)} DESC
            """
        )
        grouped: dict[datetime, dict[str, Any]] = {}
        for reading_time, result_timestamp, result_id, value in cursor.fetchall():
            result_key = str(result_id or "").strip().upper()
            if result_key not in {"R1", "R2"} or value is None:
                continue
            ts = coerce_datetime(result_timestamp) or coerce_datetime(reading_time)
            if ts is None:
                continue
            try:
                value_number = round(float(value), 6)
            except (TypeError, ValueError):
                continue
            entry = grouped.setdefault(
                ts,
                {
                    "timestamp": ts,
                    "r1": None,
                    "r2": None,
                    "sequence": state.sequence,
                    "program_a": state.program_a,
                    "program_b": state.program_b,
                },
            )
            if result_key == "R1":
                entry["r1"] = value_number
            elif result_key == "R2":
                entry["r2"] = value_number

        rows = [
            row
            for _, row in sorted(grouped.items(), key=lambda item: item[0])
            if row["r1"] is not None or row["r2"] is not None
        ]
        state.connection_ok = True
        return rows[-12:]
    except Exception as exc:
        app.logger.warning("Could not load titrator history from SQL Server: %s", exc)
        state.connection_ok = False
        return fallback
    finally:
        connection.close()


def clear_titrator_storage() -> bool:
    if not titrator_sql_source_enabled():
        history.clear()
        titrator_notes.clear()
        save_titrator_notes()
        command_log.clear()
        save_command_log()
        return True

    connection = get_sql_connection()
    if connection is None:
        return False

    try:
        ensure_titrator_table(connection)
        cursor = connection.cursor()
        cursor.execute(f"DELETE FROM {sql_table_name(SQL_TITRATOR_TABLE)}")
        connection.commit()
        history.clear()
        titrator_notes.clear()
        save_titrator_notes()
        command_log.clear()
        save_command_log()
        state.connection_ok = True
        return True
    except Exception as exc:
        app.logger.warning("Could not clear titrator SQL table: %s", exc)
        state.connection_ok = False
        return False
    finally:
        connection.close()


def save_titrator_manual_reading(timestamp: datetime, r1: float | None, r2: float | None, note: str = "") -> bool:
    if not titrator_sql_source_enabled():
        history.append(
            {
                "timestamp": timestamp,
                "r1": round(r1, 6) if r1 is not None else None,
                "r2": round(r2, 6) if r2 is not None else None,
                "sequence": state.sequence,
                "program_a": state.program_a,
                "program_b": state.program_b,
            }
        )
        del history[:-12]
        if note.strip():
            set_titrator_note(timestamp, note)
        return True

    columns = [
        SQL_TITRATOR_READING_TIME_COLUMN,
        SQL_TITRATOR_TIMESTAMP_COLUMN,
        SQL_TITRATOR_RESULT_ID_COLUMN,
        SQL_TITRATOR_VALUE_COLUMN,
    ]
    placeholders = ", ".join("?" for _ in columns)
    try:
        connection = get_sql_connection()
        if connection is None:
            raise RuntimeError("SQL connection unavailable")
        ensure_titrator_table(connection)
        cursor = connection.cursor()
        for result_id, value in [("R1", r1), ("R2", r2)]:
            if value is None:
                continue
            cursor.execute(
                f"""
                INSERT INTO {sql_table_name(SQL_TITRATOR_TABLE)} ({sql_column_list(columns)})
                VALUES ({placeholders})
                """,
                timestamp,
                timestamp,
                result_id,
                value,
            )
        connection.commit()
        state.connection_ok = True
        if note.strip():
            set_titrator_note(timestamp, note)
        return True
    except Exception as exc:
        app.logger.warning("Could not save titrator manual reading to SQL Server: %s", exc)
        state.connection_ok = False
        return False
    finally:
        if "connection" in locals() and connection is not None:
            connection.close()


def load_voltage_entries() -> list[dict[str, Any]]:
    global _voltage_channel_cols
    fallback = list(voltage_entries)
    if not sql_server_enabled():
        return fallback

    connection = get_sql_connection()
    if connection is None:
        return fallback

    try:
        resolved_table = resolve_voltage_table_name(connection)
        if not resolved_table:
            _voltage_channel_cols = []
            return []
        channel_cols = discover_voltage_channel_columns(connection, resolved_table)
        if not channel_cols:
            _voltage_channel_cols = []
            return []
        channel_cols = filter_voltage_channels_with_data(connection, resolved_table, channel_cols)
        if not channel_cols:
            _voltage_channel_cols = []
            return []
        _voltage_channel_cols = channel_cols
        table_column_names = discover_table_columns(connection, SQL_SCHEMA, resolved_table)
        table_columns = {column.lower() for column in table_column_names}
        timestamp_column = choose_voltage_timestamp_column(table_column_names)
        if not timestamp_column:
            app.logger.warning("Could not find a timestamp column in voltage table %s.", resolved_table)
            return fallback
        select_columns = [timestamp_column] + channel_cols
        include_note = voltage_note_sql_enabled() and SQL_VOLTAGE_NOTE_COLUMN.lower() in table_columns
        if include_note:
            select_columns.append(SQL_VOLTAGE_NOTE_COLUMN)
        order_columns = [sql_identifier(timestamp_column) + " DESC"]
        if SQL_VOLTAGE_ROW_ID_COLUMN and SQL_VOLTAGE_ROW_ID_COLUMN.lower() in table_columns:
            order_columns.append(sql_identifier(SQL_VOLTAGE_ROW_ID_COLUMN) + " DESC")
        cursor = connection.cursor()
        cursor.execute(
            f"""
            SELECT TOP (200)
                {sql_column_list(select_columns)}
            FROM {sql_table_name(resolved_table)}
            ORDER BY {", ".join(order_columns)}
            """
        )
        rows: list[dict[str, Any]] = []
        for db_row in cursor.fetchall():
            measured_at = db_row[0]
            note = db_row[len(channel_cols) + 1] if include_note and len(db_row) > len(channel_cols) + 1 else ""
            channels = {col: text_or_default(db_row[i + 1], "0") for i, col in enumerate(channel_cols)}
            ts = coerce_datetime(measured_at)
            if ts is None:
                continue
            ts_key = ts.isoformat(timespec="seconds")
            rows.append(
                {
                    "timestamp": format_timestamp(ts),
                    "ts_raw": ts,
                    "timestamp_key": ts_key,
                    "channels": channels,
                    "note": text_or_default(note, "") or voltage_notes.get(ts_key, ""),
                }
            )
        return list(reversed(rows))
    except Exception as exc:
        app.logger.warning("Could not load voltage entries from SQL Server: %s", exc)
        return fallback
    finally:
        connection.close()


def save_voltage_entry(timestamp: datetime, channels: dict[str, str], note: str) -> None:
    if not sql_server_enabled():
        voltage_entries.append(
            {
                "timestamp": format_timestamp(timestamp),
                "ts_raw": timestamp,
                "timestamp_key": timestamp.isoformat(timespec="seconds"),
                "channels": channels,
                "note": note,
            }
        )
        return

    connection = None
    try:
        connection = get_sql_connection()
        if connection is None:
            raise RuntimeError("SQL connection unavailable")
        resolved_table = resolve_voltage_table_name(connection)
        if not resolved_table:
            raise RuntimeError("No voltage SQL table could be resolved")
        table_column_names = discover_table_columns(connection, SQL_SCHEMA, resolved_table)
        table_columns = {column.lower(): column for column in table_column_names}
        timestamp_column = choose_voltage_timestamp_column(table_column_names) or SQL_VOLTAGE_TIMESTAMP_COLUMN
        insert_channels = {
            table_columns.get(column.lower(), column): value
            for column, value in channels.items()
            if column.lower() in table_columns
        }
        columns = [timestamp_column] + list(insert_channels.keys())
        values: list[Any] = [timestamp] + [_safe_float(v) for v in insert_channels.values()]
        if voltage_note_sql_enabled() and SQL_VOLTAGE_NOTE_COLUMN.lower() in table_columns:
            columns.append(table_columns[SQL_VOLTAGE_NOTE_COLUMN.lower()])
            values.append(note or None)
        placeholders = ", ".join("?" for _ in columns)
        cursor = connection.cursor()
        cursor.execute(
            f"""
            INSERT INTO {sql_table_name(resolved_table)} ({sql_column_list(columns)})
            VALUES ({placeholders})
            """,
            *values,
        )
        connection.commit()
    except Exception as exc:
        app.logger.warning("Could not save voltage entry to SQL Server: %s", exc)
        voltage_entries.append(
            {
                "timestamp": format_timestamp(timestamp),
                "ts_raw": timestamp,
                "timestamp_key": timestamp.isoformat(timespec="seconds"),
                "channels": channels,
                "note": note,
            }
        )
    finally:
        if "connection" in locals() and connection is not None:
            connection.close()


def refresh_chlorine_cache() -> list[dict[str, Any]]:
    rows = load_chlorine_measurements()
    chlorine_measurements.clear()
    chlorine_measurements.extend(rows)
    return list(chlorine_measurements)


def refresh_results_cache() -> list[dict[str, Any]]:
    rows = load_results_entries()
    results_entries.clear()
    results_entries.extend(rows)
    return list(results_entries)


def refresh_titrator_cache() -> list[dict[str, Any]]:
    if titrator_sql_source_enabled():
        rows = load_titrator_history()
        history.clear()
        history.extend(rows)
    return list(history)


def refresh_voltage_cache() -> list[dict[str, Any]]:
    rows = load_voltage_entries()
    voltage_entries.clear()
    voltage_entries.extend(rows)
    return list(voltage_entries)


def handle_command_state(command: str) -> None:
    if command == "Start Sequence":
        state.sequence = "Running"
    elif command == "Stop Sequence":
        state.sequence = "Stopped"
    elif command == "Break Sequence":
        state.sequence = "Paused"
    elif command == "Trigger Program A":
        state.program_a = "Active"
    elif command == "Trigger Program B":
        state.program_b = "Active"
    elif command == "Trigger Program X":
        state.sequence = "Running"


def bootstrap_data() -> None:
    global _bootstrapped
    if _bootstrapped:
        return
    _bootstrapped = True
    load_titrator_notes()
    load_voltage_notes()
    load_command_logs()

    if titrator_sql_source_enabled():
        refresh_titrator_cache()
    else:
        history.extend(create_seeded_reading(index) for index in range(12))

    if history:
        state.sequence = history[-1]["sequence"]
        state.program_a = history[-1]["program_a"]
        state.program_b = history[-1]["program_b"]

    if sql_server_enabled():
        refresh_results_cache()
        refresh_voltage_cache()
        refresh_chlorine_cache()
    else:
        _ct = [datetime.now() - timedelta(hours=2), datetime.now() - timedelta(minutes=35)]
        chlorine_measurements.extend(
            [
                {
                    "timestamp": _ct[0],
                    "timestamp_key": _ct[0].isoformat(timespec="seconds"),
                    "value": 2.15,
                    "unit": "mg/L",
                    "note": "Shift start manual check",
                },
                {
                    "timestamp": _ct[1],
                    "timestamp_key": _ct[1].isoformat(timespec="seconds"),
                    "value": 2.32,
                    "unit": "mg/L",
                    "note": "",
                },
            ]
        )
        results_entries.extend(seed_results_entries(datetime.now()))

    if not command_log:
        command_log.extend(
            [
                {"command": "Dashboard initialized", "address": "--", "time": datetime.now().strftime("%I:%M:%S %p")},
                {"command": "Monitoring active", "address": "--", "time": datetime.now().strftime("%I:%M:%S %p")},
            ]
        )
        save_command_log()

    if not sql_server_enabled():
        now = datetime.now()
        _vt = [now - timedelta(hours=2), now - timedelta(hours=1, minutes=30), now - timedelta(hours=1), now - timedelta(minutes=30), now - timedelta(minutes=10)]
        voltage_entries.extend(
            [
                {"timestamp": format_timestamp(_vt[0]), "ts_raw": _vt[0], "timestamp_key": _vt[0].isoformat(timespec="seconds"), "channels": {SQL_VOLTAGE_CH1_COLUMN: "3.52", SQL_VOLTAGE_CH2_COLUMN: "3.48", SQL_VOLTAGE_CH3_COLUMN: "3.55", SQL_VOLTAGE_CH4_COLUMN: "3.50"}, "note": "Shift start"},
                {"timestamp": format_timestamp(_vt[1]), "ts_raw": _vt[1], "timestamp_key": _vt[1].isoformat(timespec="seconds"), "channels": {SQL_VOLTAGE_CH1_COLUMN: "3.54", SQL_VOLTAGE_CH2_COLUMN: "3.51", SQL_VOLTAGE_CH3_COLUMN: "3.57", SQL_VOLTAGE_CH4_COLUMN: "3.53"}, "note": ""},
                {"timestamp": format_timestamp(_vt[2]), "ts_raw": _vt[2], "timestamp_key": _vt[2].isoformat(timespec="seconds"), "channels": {SQL_VOLTAGE_CH1_COLUMN: "3.57", SQL_VOLTAGE_CH2_COLUMN: "3.53", SQL_VOLTAGE_CH3_COLUMN: "3.59", SQL_VOLTAGE_CH4_COLUMN: "3.55"}, "note": ""},
                {"timestamp": format_timestamp(_vt[3]), "ts_raw": _vt[3], "timestamp_key": _vt[3].isoformat(timespec="seconds"), "channels": {SQL_VOLTAGE_CH1_COLUMN: "3.59", SQL_VOLTAGE_CH2_COLUMN: "3.55", SQL_VOLTAGE_CH3_COLUMN: "3.61", SQL_VOLTAGE_CH4_COLUMN: "3.57"}, "note": ""},
                {"timestamp": format_timestamp(_vt[4]), "ts_raw": _vt[4], "timestamp_key": _vt[4].isoformat(timespec="seconds"), "channels": {SQL_VOLTAGE_CH1_COLUMN: "3.60", SQL_VOLTAGE_CH2_COLUMN: "3.56", SQL_VOLTAGE_CH3_COLUMN: "3.62", SQL_VOLTAGE_CH4_COLUMN: "3.58"}, "note": ""},
            ]
        )


@app.get("/")
def home_page() -> str:
    bootstrap_data()
    chlorine_rows = serialize_chlorine(refresh_chlorine_cache())
    result_rows = refresh_results_cache()
    titrator_rows = refresh_titrator_cache()
    voltage_rows = refresh_voltage_cache()
    wincc_sources = load_wincc_sources()
    latest_chlorine = chlorine_rows[0] if chlorine_rows else None
    latest_results = get_results_rows()[0] if result_rows else None
    latest_voltage = list(reversed(voltage_rows))[0] if voltage_rows else None
    serialized_history = serialize_history(titrator_rows)
    latest_titrator = serialized_history[-1] if serialized_history else None
    return render_template(
        "home.html",
        active_tab="home",
        announcements=list(reversed(announcements)),
        latest_chlorine=latest_chlorine,
        latest_results=latest_results,
        latest_voltage=latest_voltage,
        latest_titrator=latest_titrator,
        titrator_summary=latest_summary(),
        titrator_history=serialized_history,
        chlorine_measurements=chlorine_rows,
        voltage_entries=list(reversed(voltage_rows)),
        voltage_channel_columns=_voltage_channel_cols,
        wincc_sources=wincc_sources,
    )


@app.post("/announcement")
def create_announcement() -> Any:
    global _announcement_id_counter
    message = (request.form.get("message") or "").strip()
    author = (request.form.get("author") or "Team").strip() or "Team"
    if message:
        _announcement_id_counter += 1
        announcements.append({
            "id": _announcement_id_counter,
            "timestamp": format_timestamp(datetime.now()),
            "author": author,
            "message": message,
        })
        flash("Announcement posted.")
    return redirect(url_for("home_page"))


@app.post("/announcement/<int:ann_id>/delete")
def delete_announcement(ann_id: int) -> Any:
    global announcements
    announcements = [a for a in announcements if a["id"] != ann_id]
    return redirect(url_for("home_page"))


@app.post("/clear/results")
def clear_results() -> Any:
    if clear_results_storage():
        refresh_results_cache()
        flash("Results table cleared.")
    else:
        flash("Results table could not be cleared.")
    return redirect(url_for("results_page"))


@app.post("/clear/chlorine")
def clear_chlorine() -> Any:
    if clear_chlorine_storage():
        refresh_chlorine_cache()
        flash("Chlorine table cleared.")
    else:
        flash("Chlorine table could not be cleared.")
    return redirect(url_for("chlorine_page"))


@app.post("/clear/history")
def clear_history() -> Any:
    if clear_titrator_storage():
        refresh_titrator_cache()
        flash("Titrator table cleared.")
    else:
        flash("Titrator table could not be cleared.")
    return redirect(url_for("titrator_page"))


@app.post("/clear/voltage")
def clear_voltage() -> Any:
    if sql_server_enabled():
        connection = get_sql_connection()
        if connection:
            try:
                resolved_table = resolve_voltage_table_name(connection)
                if resolved_table:
                    cursor = connection.cursor()
                    cursor.execute(f"DELETE FROM {sql_table_name(resolved_table)}")
                    connection.commit()
                    voltage_entries.clear()
                    voltage_notes.clear()
                    save_voltage_notes()
                    voltage_command_log.clear()
                    save_voltage_command_log()
                    flash("Voltage table cleared.")
                else:
                    flash("Voltage table could not be found.")
            except Exception as exc:
                app.logger.warning("Could not clear voltage SQL table: %s", exc)
                flash("Voltage table could not be cleared.")
            finally:
                connection.close()
        else:
            flash("No SQL connection available.")
    else:
        voltage_entries.clear()
        voltage_notes.clear()
        save_voltage_notes()
        voltage_command_log.clear()
        save_voltage_command_log()
        flash("Voltage table cleared.")
    return redirect(url_for("voltage_page"))


@app.post("/clear/wincc")
def clear_wincc() -> Any:
    table_ref = (request.form.get("table_ref") or "").strip()
    ok, message = clear_wincc_table(table_ref)
    flash(message if ok else f"WinCC table could not be cleared. {message}")
    return redirect(url_for("wincc_page"))


@app.post("/clear/all")
def clear_all() -> Any:
    bootstrap_data()
    tables = request.form.getlist("tables")
    if not tables:
        flash("No tables selected.")
        return redirect(url_for("home_page"))
    cleared: list[str] = []
    failed: list[str] = []
    if "titrator" in tables:
        if clear_titrator_storage():
            refresh_titrator_cache()
            cleared.append("Titrator")
        else:
            failed.append("Titrator")
    if "chlorine" in tables:
        if clear_chlorine_storage():
            refresh_chlorine_cache()
            cleared.append("Chlorine")
        else:
            failed.append("Chlorine")
    if "results" in tables:
        if clear_results_storage():
            refresh_results_cache()
            cleared.append("Results")
        else:
            failed.append("Results")
    if "voltage" in tables:
        if sql_server_enabled():
            connection = get_sql_connection()
            if connection:
                try:
                    resolved_table = resolve_voltage_table_name(connection)
                    if resolved_table:
                        cursor = connection.cursor()
                        cursor.execute(f"DELETE FROM {sql_table_name(resolved_table)}")
                        connection.commit()
                        voltage_entries.clear()
                        voltage_notes.clear()
                        save_voltage_notes()
                        voltage_command_log.clear()
                        save_voltage_command_log()
                        cleared.append("Voltage")
                    else:
                        failed.append("Voltage")
                except Exception as exc:
                    app.logger.warning("Could not clear voltage in clear_all: %s", exc)
                    failed.append("Voltage")
                finally:
                    connection.close()
            else:
                failed.append("Voltage")
        else:
            voltage_entries.clear()
            voltage_notes.clear()
            save_voltage_notes()
            voltage_command_log.clear()
            save_voltage_command_log()
            cleared.append("Voltage")
    if "wincc" in tables:
        wincc_cleared = []
        wincc_failed = []
        for table_ref in SQL_WINCC_TABLE_NAMES:
            ok, _ = clear_wincc_table(table_ref)
            if ok:
                wincc_cleared.append(table_ref)
            else:
                wincc_failed.append(table_ref)
        if wincc_cleared:
            cleared.append("WinCC")
        if wincc_failed:
            failed.append("WinCC")
    if cleared:
        flash(f"Cleared: {', '.join(cleared)}.")
    if failed:
        flash(f"Could not clear: {', '.join(failed)}.")
    return redirect(url_for("home_page"))


@app.post("/export/all.xlsx")
def export_all() -> Any:
    bootstrap_data()
    tables = request.form.getlist("tables")
    if not tables:
        flash("No tables selected.")
        return redirect(url_for("home_page"))

    bold = Font(bold=True)
    center = Alignment(horizontal="center", vertical="center")
    wb = Workbook()
    first_sheet = True

    def _sheet(ws: Any, headers: list[str], rows: list[list[Any]]) -> None:
        for col_idx, h in enumerate(headers, start=1):
            c = ws.cell(row=1, column=col_idx, value=h)
            c.font = bold
            c.alignment = center
        for row_idx, row in enumerate(rows, start=2):
            for col_idx, val in enumerate(row, start=1):
                ws.cell(row=row_idx, column=col_idx, value=val)
        for col in ws.columns:
            max_len = max((len(str(c.value or "")) for c in col), default=8)
            ws.column_dimensions[get_column_letter(col[0].column)].width = min(max_len + 4, 30)

    def _add_sheet(title: str, headers: list[str], rows: list[list[Any]]) -> None:
        nonlocal first_sheet
        if first_sheet:
            ws = wb.active
            ws.title = title
            first_sheet = False
        else:
            ws = wb.create_sheet(title)
        _sheet(ws, headers, rows)

    if "results" in tables:
        result_rows = refresh_results_cache()
        _add_sheet("Results",
            ["Timestamp", "Time Running (hrs)", "Time Flushing (hrs)", "Cell to Cell (V)",
             "Bus 2 Bus (V)", "Tank Temp (C)", "Catholyte pH", "Anolyte pH",
             "Catholyte Flow (mL/min)", "Anolyte Flow (mL/min)", "Stack Current (A)",
             "Current Density (A/m2)", "Ca Precipitated", "Mg Precipitated",
             "OH- in Solution", "Total OH-", "Theoretical OH-",
             "Catholyte Dosage (C/mL)", "Anolyte Dosage (C/mL)",
             "CE OH- Titration (%)", "CIER", "Note"],
            [[export_timestamp(r["ts_raw"]) if isinstance(r.get("ts_raw"), datetime) else r["timestamp"],
              r.get("time_running",""), r.get("time_flushing",""), r.get("cell_to_cell",""),
              r.get("bus_2_bus",""), r.get("tank_temp",""), r.get("catholyte_ph",""), r.get("anolyte_ph",""),
              r.get("catholyte_flow",""), r.get("anolyte_flow",""), r.get("stack_current",""),
              r.get("current_density",""), r.get("ca_precipitated",""), r.get("mg_precipitated",""),
              r.get("oh_solution",""), r.get("total_oh",""), r.get("theoretical_oh",""),
              r.get("catholyte_dosage",""), r.get("anolyte_dosage",""),
              r.get("ce_oh_titration",""), r.get("cier",""), r.get("note","")]
             for r in result_rows])

    if "chlorine" in tables:
        chlorine_rows = refresh_chlorine_cache()
        _add_sheet("Chlorine",
            ["Timestamp", "Measurement (mg/L)", "Note"],
            [[export_timestamp(r["timestamp"]) if isinstance(r["timestamp"], datetime) else r["timestamp"],
              f'{r["value"]:.2f}', r.get("note") or ""]
             for r in reversed(chlorine_rows)])

    if "titrator" in tables:
        titrator_source_rows = refresh_titrator_cache()
        titrator_rows = []
        for r in titrator_source_rows:
            ts = export_timestamp(r["timestamp"])
            note = get_titrator_note(r["timestamp"])
            if r["r1"] is not None:
                titrator_rows.append([ts, "R1", format_titrator_value(r["r1"]), note])
            if r["r2"] is not None:
                titrator_rows.append([ts, "R2", format_titrator_value(r["r2"]), note])
        _add_sheet("Titrator", ["Timestamp", "Result_ID", "Value", "Note"], titrator_rows)

    if "voltage" in tables:
        voltage_rows = refresh_voltage_cache()
        _vcols = _voltage_channel_cols
        _vheaders = [col.replace("_", " ").title() + " (V)" for col in _vcols]
        _add_sheet("Voltage",
            ["Timestamp"] + _vheaders + ["Note"],
            [[export_timestamp(r["ts_raw"]) if isinstance(r.get("ts_raw"), datetime) else r["timestamp"],
              *[r.get("channels", {}).get(col, "") for col in _vcols],
              r.get("note", "")]
             for r in voltage_rows])

    if "wincc" in tables:
        used_titles: set[str] = set()
        for source in load_wincc_sources():
            ws = wb.active if first_sheet else wb.create_sheet()
            if first_sheet:
                first_sheet = False
            ws.title = safe_worksheet_title(source["table"], used_titles)
            headers = [column["label"] for column in source.get("columns", [])]
            rows = source.get("rows", [])
            if headers:
                ws.append(headers)
                for row in reversed(rows):
                    ws.append(row)
            else:
                ws.append(["Message"])
                ws.append([source.get("error") or "No rows available."])
            for cell in ws[1]:
                cell.font = Font(bold=True)
                cell.alignment = Alignment(horizontal="center")
            for col in ws.columns:
                max_len = max((len(str(c.value or "")) for c in col), default=8)
                ws.column_dimensions[get_column_letter(col[0].column)].width = min(max_len + 4, 30)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    sheet_names = [s.title for s in wb.worksheets]
    filename = "-".join(n.lower() for n in sheet_names) + ".xlsx"
    return send_file(output, mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                     as_attachment=True, download_name=filename)


@app.get("/results")
def results_page() -> str:
    bootstrap_data()
    refresh_results_cache()
    voltage_rows = refresh_voltage_cache()
    return render_template(
        "results.html",
        active_tab="results",
        result_cards=get_result_cards(),
        result_rows=get_results_rows(),
        voltage_entries=list(reversed(voltage_rows)),
        voltage_channel_columns=_voltage_channel_cols,
    )


@app.post("/results")
def create_results_entry() -> Any:
    bootstrap_data()
    f = request.form
    ts_text = (f.get("timestamp") or "").strip()
    try:
        ts = datetime.fromisoformat(ts_text) if ts_text else datetime.now()
    except ValueError:
        ts = datetime.now()

    def val(key: str) -> str:
        return (f.get(key) or "—").strip() or "—"

    entry = {
        "timestamp": format_timestamp(ts), "ts_raw": ts, "timestamp_key": ts.isoformat(timespec="seconds"),
        "time_running": val("time_running"), "time_flushing": val("time_flushing"),
        "cell_to_cell": val("cell_to_cell"), "bus_2_bus": val("bus_2_bus"), "tank_temp": val("tank_temp"),
        "catholyte_ph": val("catholyte_ph"), "anolyte_ph": val("anolyte_ph"),
        "catholyte_flow": val("catholyte_flow"), "anolyte_flow": val("anolyte_flow"),
        "stack_current": val("stack_current"), "current_density": val("current_density"),
        "ca_precipitated": val("ca_precipitated"), "mg_precipitated": val("mg_precipitated"),
        "oh_solution": val("oh_solution"), "total_oh": val("total_oh"), "theoretical_oh": val("theoretical_oh"),
        "catholyte_dosage": val("catholyte_dosage"), "anolyte_dosage": val("anolyte_dosage"),
        "ce_oh_titration": val("ce_oh_titration"), "cier": val("cier"),
        "note": (f.get("note") or "").strip(),
    }
    save_results_entry(entry)
    refresh_results_cache()
    flash("Results entry saved.")
    return redirect(url_for("results_page"))


@app.post("/voltage")
def create_voltage_entry() -> Any:
    bootstrap_data()
    f = request.form
    ts_text = (f.get("timestamp") or "").strip()
    try:
        ts = datetime.fromisoformat(ts_text) if ts_text else datetime.now()
    except ValueError:
        ts = datetime.now()
    excluded = {"timestamp", "note"}
    channels = {k: (v or "0").strip() or "0" for k, v in f.items() if k not in excluded}
    if not channels:
        channels = {col: "0" for col in _voltage_channel_cols}
    save_voltage_entry(ts, channels, (f.get("note") or "").strip())
    refresh_voltage_cache()
    flash("Voltage entry saved.")
    return redirect(url_for("voltage_page"))


@app.get("/chlorine")
def chlorine_page() -> str:
    bootstrap_data()
    chlorine_rows = refresh_chlorine_cache()
    return render_template(
        "chlorine.html",
        page_title="Chlorine Measurements",
        hero_title="Chlorine entries tracked separately from the rest of the machine data.",
        hero_copy="This tab is for entering and reviewing chlorine measurements only, with timestamp and notes saved in their own flow.",
        active_tab="chlorine",
        connection_text="Connected to SQL-ready local app",
        chlorine_measurements=serialize_chlorine(chlorine_rows),
    )


@app.get("/titrator")
def titrator_page() -> str:
    bootstrap_data()
    titrator_rows = refresh_titrator_cache()
    return render_template(
        "titrator.html",
        page_title="Titrator Dashboard",
        hero_title="Titrator readings, controls, and R1/R2 history in one place.",
        hero_copy="This tab keeps the instrument-facing data together so it does not get mixed with chlorine entry or final calculation results.",
        active_tab="titrator",
        connection_text=latest_summary()["connection_text"],
        summary=latest_summary(),
        history=serialize_history(titrator_rows),
        command_log=command_log,
        control_actions=machine_control_actions(),
    )


@app.post("/titrator/app/start")
def start_titrator_app_route() -> Any:
    bootstrap_data()
    ok, message = start_titrolyzer_app()
    if ok:
        add_command("Start Titrolyzer app", "local app")
    flash(message)
    return redirect(url_for("titrator_page"))


@app.post("/titrator/app/stop")
def stop_titrator_app_route() -> Any:
    bootstrap_data()
    ok, message = stop_titrolyzer_app()
    if ok:
        add_command("Stop Titrolyzer app", "local app")
    flash(message)
    return redirect(url_for("titrator_page"))



@app.get("/api/titrator")
def api_titrator() -> Any:
    bootstrap_data()
    titrator_rows = refresh_titrator_cache()
    data = []
    for row in titrator_rows:
        ts = export_timestamp(row["timestamp"])
        if row["r1"] is not None:
            data.append({"Timestamp": ts, "Result_ID": "R1", "Value": row["r1"]})
        if row["r2"] is not None:
            data.append({"Timestamp": ts, "Result_ID": "R2", "Value": row["r2"]})
    return jsonify(data)


@app.get("/wincc")
def wincc_page() -> str:
    bootstrap_data()
    return render_template(
        "wincc.html",
        page_title="WinCC Data",
        active_tab="wincc",
        wincc_sources=load_wincc_sources(),
        row_limit=SQL_WINCC_ROW_LIMIT,
    )


@app.get("/voltage")
def voltage_page() -> str:
    bootstrap_data()
    voltage_rows = refresh_voltage_cache()
    windaq_status = windaq_mssql_status()
    return render_template(
        "voltage.html",
        active_tab="voltage",
        voltage_entries=list(reversed(voltage_rows)),
        voltage_command_log=voltage_command_log,
        voltage_channel_columns=_voltage_channel_cols,
        windaq_running=windaq_status["running"],
        windaq_status=windaq_status["label"],
    )


@app.get("/api/windaq-status")
def api_windaq_status() -> Any:
    bootstrap_data()
    status = windaq_mssql_status()
    return jsonify({"running": status["running"], "label": status["label"]})


@app.get("/api/windaq-debug")
def api_windaq_debug() -> Any:
    """Lists all child controls of the WinDAQ-MSSQL window for diagnostics."""
    if os.name != "nt":
        return jsonify({"error": "Windows only"})
    hwnd = _win32_find_windaq_window(timeout_seconds=2.0)
    if not hwnd:
        return jsonify({"error": "WinDAQ-MSSQL window not found"})
    controls = []
    for child in _win32_enum_child_windows(hwnd):
        import ctypes
        user32 = ctypes.windll.user32
        controls.append({
            "hwnd": child,
            "class": _win32_class_name(child),
            "text": _win32_text(child),
            "visible": bool(user32.IsWindowVisible(child)),
            "enabled": bool(user32.IsWindowEnabled(child)),
        })
    return jsonify({"hwnd": hwnd, "controls": controls})


@app.post("/voltage/command")
def voltage_command() -> Any:
    command = (request.form.get("command") or "").strip()
    message = f"{command} logged."
    if command == "Start WinDAQ-MSSQL":
        _, message = start_windaq_mssql()
    elif command == "Stop WinDAQ-MSSQL":
        _, message = stop_windaq_mssql()
    if command:
        voltage_command_log.append({
            "command": command,
            "time": datetime.now().strftime("%I:%M:%S %p"),
        })
        save_voltage_command_log()
    flash(message)
    return redirect(url_for("voltage_page"))


@app.get("/export/results.xlsx")
def export_results() -> Any:
    bootstrap_data()
    rows = refresh_results_cache()

    wb = Workbook()
    ws = wb.active
    ws.title = "Results"

    green_fill = PatternFill("solid", fgColor="4CAF50")
    bold = Font(bold=True)
    center = Alignment(horizontal="center", vertical="center")

    # Row 1 — group header
    ws.merge_cells("B1:L1")
    cell = ws["B1"]
    cell.value = "Experiment inputs"
    cell.fill = green_fill
    cell.font = Font(bold=True, color="FFFFFF")
    cell.alignment = center

    # Row 2 — column names with units
    col_names = [
        "Timestamp", "Time Running (Hours)", "Time Flushing (Hours)",
        "Cell to cell (V)", "Bus 2 Bus (V)", "Tank temp (C)",
        "Catholyte pH", "Anolyte pH",
        "Catholyte flow rate (mL/min)", "Anolyte flow rate (mL/min)",
        "Stack Current (A)", "Current Density (A/m2)",
        "Ca precipitated (mol/sec)", "Mg precipitated (mol/sec)",
        "OH- produced in solution (mol/sec)", "Total OH- produced (mol/sec)",
        "Theoretical OH- produced (mol/sec)",
        "Catholyte dosage (C/mL)", "Anolyte dosage (C/mL)",
        "CE OH- (TITRATION) (%)", "CIER", "Note",
    ]
    for col_idx, name in enumerate(col_names, start=1):
        c = ws.cell(row=2, column=col_idx, value=name)
        c.font = bold
        c.alignment = center

    # Data rows starting at row 3
    fields = [
        "ts_raw", "time_running", "time_flushing", "cell_to_cell", "bus_2_bus",
        "tank_temp", "catholyte_ph", "anolyte_ph", "catholyte_flow", "anolyte_flow",
        "stack_current", "current_density", "ca_precipitated", "mg_precipitated",
        "oh_solution", "total_oh", "theoretical_oh",
        "catholyte_dosage", "anolyte_dosage", "ce_oh_titration", "cier",
    ]
    for row_idx, row in enumerate(rows, start=3):
        for col_idx, field in enumerate(fields, start=1):
            val = row.get(field, "")
            if field == "ts_raw" and isinstance(val, datetime):
                val = export_timestamp(val)
            ws.cell(row=row_idx, column=col_idx, value=val)
        ws.cell(row=row_idx, column=len(fields) + 1, value=row.get("note", ""))

    for col in ws.columns:
        max_len = max((len(str(c.value or "")) for c in col), default=8)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(max_len + 4, 30)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return send_file(output, mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                     as_attachment=True, download_name="results-table.xlsx")


@app.get("/export/chlorine.xlsx")
def export_chlorine() -> Any:
    bootstrap_data()
    chlorine_rows = refresh_chlorine_cache()
    output = simple_xlsx(
        ["Timestamp", "Measurement (mg/L)", "Note"],
        [[export_timestamp(r["timestamp"]) if isinstance(r["timestamp"], datetime) else r["timestamp"],
          f'{r["value"]:.2f}', r.get("note") or ""] for r in reversed(chlorine_rows)],
        "Chlorine"
    )
    return send_file(output, mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                     as_attachment=True, download_name="chlorine-table.xlsx")


@app.get("/api/dashboard")
def api_dashboard() -> Any:
    bootstrap_data()
    if request.args.get("refresh") == "1":
        append_reading()
    titrator_rows = refresh_titrator_cache()
    chlorine_rows = refresh_chlorine_cache()

    return jsonify(
        {
            "summary": latest_summary(),
            "history": serialize_history(titrator_rows),
            "command_log": command_log,
            "chlorine_measurements": serialize_chlorine(chlorine_rows),
        }
    )


@app.get("/api/voltage")
def api_voltage() -> Any:
    bootstrap_data()
    rows = refresh_voltage_cache()
    return jsonify([
        {"chart_time": r["ts_raw"].strftime("%H:%M") if isinstance(r.get("ts_raw"), datetime) else r["timestamp"],
         "channels": r.get("channels", {})}
        for r in rows
    ])


@app.post("/api/command")
def api_command() -> Any:
    bootstrap_data()
    payload = request.get_json(silent=True) or {}
    command = str(payload.get("command", "")).strip()
    address = str(payload.get("address", "")).strip()

    if not command or not address:
        return jsonify({"ok": False, "error": "Missing command or address."}), 400
    expected_address = COMMAND_COIL_MANUAL_ADDRESS_MAP.get(command)
    if expected_address is None:
        return jsonify({"ok": False, "error": "Unknown machine command."}), 400
    if address != str(expected_address):
        return jsonify({"ok": False, "error": "Command address does not match the manual coil map."}), 400

    handle_command_state(command)
    add_command(command, address)
    command_sent = write_modbus_coil(command)
    append_reading()
    if modbus_enabled() and not command_sent:
        return jsonify({"ok": False, "error": f"{command} could not be written to coil {address}."}), 502
    return jsonify(
        {
            "ok": True,
            "message": f"{command} sent to manual coil address {address}.",
            "summary": latest_summary(),
            "history": serialize_history(refresh_titrator_cache()),
            "command_log": command_log,
        }
    )


@app.post("/api/manual-reading")
def api_manual_reading() -> Any:
    bootstrap_data()
    payload = request.get_json(silent=True) or {}
    try:
        r1 = float(payload.get("r1", 0))
        r2 = float(payload.get("r2", 0))
    except (TypeError, ValueError):
        return jsonify({"ok": False, "error": "Invalid values."}), 400
    note = str(payload.get("note", "")).strip()
    saved = save_titrator_manual_reading(datetime.now(), r1, r2, note)
    if not saved:
        return jsonify({"ok": False, "error": "Could not save titrator reading to SQL Server."}), 500
    titrator_rows = refresh_titrator_cache()
    return jsonify({"ok": True, "summary": latest_summary(), "history": serialize_history(titrator_rows), "command_log": command_log})


@app.post("/api/titrator-note")
def api_titrator_note() -> Any:
    bootstrap_data()
    payload = request.get_json(silent=True) or {}
    timestamp = parse_titrator_timestamp(str(payload.get("timestamp", "")))
    if timestamp is None:
        return jsonify({"ok": False, "error": "Missing or invalid timestamp."}), 400
    note = str(payload.get("note", "")).strip()
    set_titrator_note(timestamp, note)
    titrator_rows = refresh_titrator_cache()
    return jsonify(
        {
            "ok": True,
            "history": serialize_history(titrator_rows),
            "saved_note": note,
        }
    )


@app.post("/api/results-note")
def api_results_note() -> Any:
    bootstrap_data()
    payload = request.get_json(silent=True) or {}
    rowkey = str(payload.get("rowkey", "")).strip()
    note = str(payload.get("note", "")).strip()
    if not rowkey:
        return jsonify({"ok": False, "error": "Missing rowkey."}), 400
    for entry in results_entries:
        if entry.get("timestamp_key") == rowkey:
            entry["note"] = note
            break
    if sql_server_enabled():
        try:
            ts = datetime.fromisoformat(rowkey)
            connection = get_sql_connection()
            if connection:
                cursor = connection.cursor()
                cursor.execute(
                    f"UPDATE {sql_table_name(SQL_RESULTS_TABLE)} SET note = ? WHERE measured_at = ?",
                    note or None, ts,
                )
                connection.commit()
                connection.close()
        except Exception as exc:
            app.logger.warning("Could not update results note in SQL: %s", exc)
    return jsonify({"ok": True})


@app.post("/api/voltage-note")
def api_voltage_note() -> Any:
    bootstrap_data()
    payload = request.get_json(silent=True) or {}
    rowkey = str(payload.get("rowkey", "")).strip()
    note = str(payload.get("note", "")).strip()
    if not rowkey:
        return jsonify({"ok": False, "error": "Missing rowkey."}), 400
    for entry in voltage_entries:
        if entry.get("timestamp_key") == rowkey:
            entry["note"] = note
            break
    if note:
        voltage_notes[rowkey] = note
    else:
        voltage_notes.pop(rowkey, None)
    save_voltage_notes()
    return jsonify({"ok": True})


@app.post("/api/chlorine-note")
def api_chlorine_note() -> Any:
    bootstrap_data()
    payload = request.get_json(silent=True) or {}
    rowkey = str(payload.get("rowkey", "")).strip()
    note = str(payload.get("note", "")).strip()
    if not rowkey:
        return jsonify({"ok": False, "error": "Missing rowkey."}), 400
    for entry in chlorine_measurements:
        if entry.get("timestamp_key") == rowkey:
            entry["note"] = note
            break
    if sql_server_enabled():
        try:
            ts = datetime.fromisoformat(rowkey)
            connection = get_sql_connection()
            if connection:
                cursor = connection.cursor()
                cursor.execute(
                    f"UPDATE {sql_table_name(SQL_CHLORINE_TABLE)} SET note = ? WHERE measured_at = ?",
                    note or None, ts,
                )
                connection.commit()
                connection.close()
        except Exception as exc:
            app.logger.warning("Could not update chlorine note in SQL: %s", exc)
    return jsonify({"ok": True})


@app.post("/chlorine")
def create_chlorine_measurement() -> Any:
    bootstrap_data()
    value_text = (request.form.get("value") or "").strip()
    timestamp_text = (request.form.get("timestamp") or "").strip()
    note = (request.form.get("note") or "").strip()
    unit = (request.form.get("unit") or "mg/L").strip() or "mg/L"

    try:
        value = float(value_text)
    except ValueError:
        return redirect(url_for("chlorine_page"))

    if timestamp_text:
        try:
            timestamp = datetime.fromisoformat(timestamp_text)
        except ValueError:
            timestamp = datetime.now()
    else:
        timestamp = datetime.now()

    save_chlorine_measurement(timestamp, value, unit, note)
    flash(f"Measurement of {value} {unit} saved successfully.")
    return redirect(url_for("chlorine_page"))


@app.post("/titrator")
def create_titrator_entry() -> Any:
    bootstrap_data()
    r1_text = (request.form.get("r1") or "").strip()
    r2_text = (request.form.get("r2") or "").strip()
    timestamp_text = (request.form.get("timestamp") or "").strip()
    note = (request.form.get("note") or "").strip()

    r1 = float(r1_text) if r1_text else None
    r2 = float(r2_text) if r2_text else None

    if r1 is None and r2 is None:
        return redirect(url_for("titrator_page"))

    if timestamp_text:
        try:
            timestamp = datetime.fromisoformat(timestamp_text)
        except ValueError:
            timestamp = datetime.now()
    else:
        timestamp = datetime.now()

    saved = save_titrator_manual_reading(timestamp, r1, r2, note)
    parts = []
    if r1 is not None:
        parts.append(f"R1: {format_titrator_value(r1)}")
    if r2 is not None:
        parts.append(f"R2: {format_titrator_value(r2)}")
    if saved:
        flash(f"Titrator reading {', '.join(parts)} saved successfully.")
    else:
        flash("Titrator reading could not be saved to SQL Server.")
    return redirect(url_for("titrator_page"))


@app.get("/export/history.xlsx")
def export_history() -> Any:
    bootstrap_data()
    titrator_rows = refresh_titrator_cache()
    rows = []
    for row in titrator_rows:
        ts = export_timestamp(row["timestamp"])
        note = get_titrator_note(row["timestamp"])
        if row["r1"] is not None:
            rows.append([ts, "R1", format_titrator_value(row["r1"]), note])
        if row["r2"] is not None:
            rows.append([ts, "R2", format_titrator_value(row["r2"]), note])
    output = simple_xlsx(["Timestamp", "Result_ID", "Value", "Note"], rows, "Titrator")
    return send_file(output, mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                     as_attachment=True, download_name="titrator-table.xlsx")


@app.get("/export/wincc.xlsx")
def export_wincc() -> Any:
    bootstrap_data()
    sources = load_wincc_sources()
    wb = Workbook()
    used_titles: set[str] = set()

    for index, source in enumerate(sources):
        ws = wb.active if index == 0 else wb.create_sheet()
        ws.title = safe_worksheet_title(source["table"], used_titles)
        headers = [column["label"] for column in source.get("columns", [])]
        rows = source.get("rows", [])
        if headers:
            ws.append(headers)
            for row in reversed(rows):
                ws.append(row)
        else:
            ws.append(["Message"])
            ws.append([source.get("error") or "No rows available."])

        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.alignment = Alignment(horizontal="center")
        for column in ws.columns:
            max_len = max((len(str(cell.value or "")) for cell in column), default=8)
            ws.column_dimensions[get_column_letter(column[0].column)].width = min(max_len + 4, 30)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return send_file(output, mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                     as_attachment=True, download_name="wincc-tables.xlsx")


@app.get("/export/voltage.xlsx")
def export_voltage() -> Any:
    bootstrap_data()
    rows = refresh_voltage_cache()
    _vcols = _voltage_channel_cols
    _vheaders = [col.replace("_", " ").title() + " (V)" for col in _vcols]
    output = simple_xlsx(
        ["Timestamp"] + _vheaders + ["Note"],
        [[export_timestamp(r["ts_raw"]) if isinstance(r.get("ts_raw"), datetime) else r["timestamp"],
          *[r.get("channels", {}).get(col, "") for col in _vcols],
          r.get("note", "")] for r in rows],
        "Voltage"
    )
    return send_file(output, mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                     as_attachment=True, download_name="voltage-table.xlsx")


@app.get("/health")
def health() -> Any:
    return jsonify({"ok": True, "time": datetime.now().isoformat()})


if __name__ == "__main__":
    bootstrap_data()
    app.run(host="0.0.0.0", port=5000, debug=True)
